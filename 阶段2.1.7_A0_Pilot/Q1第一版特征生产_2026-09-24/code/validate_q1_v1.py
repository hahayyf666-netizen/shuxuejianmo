from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook


SCHEMA_VERSION = "q1-feature-v1.0"
REQUIRED_MANIFEST = {
    "sample_key", "video_id", "clip_id", "official_text", "official_text_sha256", "source_relpath",
    "source_mp4_sha256", "source_label_xlsx_sha256", "stage1_duration_proxy_sec", "alignment_trace_sha256", "output_file",
    "output_sha256", "output_bytes", "modalities", "original_effective_duration_sec", "feature_dims_json",
    "text_sequence_length", "audio_observation_length", "video_observation_length", "alignment_mode",
    "alignment_granularity", "text_audio_correspondence", "text_av_time_mapping_status", "text_present",
    "text_content_valid", "audio_present", "audio_observation_valid", "audio_speech_valid", "video_present",
    "audio_visual_time_valid", "vision_feature_valid", "face_feature_valid", "official_word_count",
    "valid_word_time_count", "text_word_valid_count", "audio_word_valid_count", "vision_word_valid_count",
    "audio_coverage_start_sec", "audio_coverage_end_sec", "video_coverage_start_sec", "video_coverage_end_sec",
    "av_common_coverage_start_sec", "av_common_coverage_end_sec", "opensmile_config_sha256",
    "roberta_asset_identity_sha256", "config_sha256", "mediapipe_model_sha256", "process_rss_bytes", "runtime_sec", "status", "error",
}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def scalar(z: Any, name: str) -> Any:
    value = z[name]
    if value.shape != ():
        raise ValueError(f"{name} must be scalar, got {value.shape}")
    return value.item()


def expected_words(text: str) -> list[tuple[int, int, str]]:
    """Independent reimplementation of the documented non-whitespace token rule."""
    result: list[tuple[int, int, str]] = []
    pending_prefix: int | None = None
    for match in re.finditer(r"\S+", text, flags=re.UNICODE):
        chunk = match.group(0)
        has_content = any(ch.isalnum() or unicodedata.category(ch)[0] in {"L", "N"} for ch in chunk)
        if has_content:
            start = pending_prefix if pending_prefix is not None else match.start()
            result.append((start, match.end(), text[start:match.end()]))
            pending_prefix = None
        elif result:
            start, _, _ = result[-1]
            result[-1] = (start, match.end(), text[start:match.end()])
        else:
            pending_prefix = match.start()
    if pending_prefix is not None or not result:
        raise ValueError("official text cannot be tokenized by the frozen rule")
    return result


def read_official_labels(path: Path) -> tuple[dict[str, str], str]:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(rows)]
    needed = {"video_id", "clip_id", "text"}
    if not needed.issubset(header):
        book.close()
        raise ValueError(f"official label table lacks {sorted(needed - set(header))}")
    indexes = {name: header.index(name) for name in needed}
    labels: dict[str, str] = {}
    for row in rows:
        video, clip = row[indexes["video_id"]], row[indexes["clip_id"]]
        if isinstance(video, float) and video.is_integer():
            video = int(video)
        if isinstance(clip, float) and clip.is_integer():
            clip = int(clip)
        key = f"{str(video).strip()}$_${str(clip).strip()}"
        text = row[indexes["text"]]
        if key in labels or not isinstance(text, str):
            book.close()
            raise ValueError(f"duplicate or nontext official label for {key}")
        labels[key] = text
    book.close()
    return labels, digest(path)


def aggregate_reference(starts: np.ndarray, ends: np.ndarray, centers: np.ndarray, values: np.ndarray,
                        word_time_valid: np.ndarray, valid_frames: np.ndarray | None = None) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    n_words, dim = len(starts), int(values.shape[1])
    output = np.zeros((n_words, dim * 2), dtype=np.float32)
    valid = np.zeros(n_words, dtype=np.uint8)
    indptr = [0]
    indexes_all: list[int] = []
    for i in range(n_words):
        if not word_time_valid[i] or not np.isfinite(starts[i]) or not np.isfinite(ends[i]) or ends[i] <= starts[i]:
            indptr.append(len(indexes_all))
            continue
        indexes = np.flatnonzero((centers >= starts[i]) & (centers < ends[i]))
        if valid_frames is not None:
            indexes = indexes[valid_frames[indexes].astype(bool)]
        indexes_all.extend(int(idx) for idx in indexes)
        indptr.append(len(indexes_all))
        if len(indexes) and np.isfinite(values[indexes]).all():
            chosen = values[indexes].astype(np.float64)
            output[i, :dim] = chosen.mean(axis=0).astype(np.float32)
            output[i, dim:] = chosen.std(axis=0, ddof=0).astype(np.float32)
            valid[i] = 1
    return output, valid, np.asarray(indptr, dtype=np.int32), np.asarray(indexes_all, dtype=np.int32)


def validate_one(root: Path, input_root: Path, row: dict[str, str], input_row: dict[str, str], labels: dict[str, str],
                 label_sha: str, audit_rows: dict[str, dict[str, str]], audit_root: Path, assets: dict[str, Any]) -> dict[str, Any]:
    key = row["sample_key"]
    if row["status"] != "PASS":
        raise ValueError(f"sample production status is {row['status']}: {row.get('error', '')}")
    for col in ("video_id", "clip_id", "official_text", "official_text_sha256", "source_relpath", "source_mp4_sha256", "source_label_xlsx_sha256"):
        if row[col] != input_row[col]:
            raise ValueError(f"manifest/input identity mismatch at {col}")
    if key not in labels or row["official_text"] != labels[key]:
        raise ValueError("NPZ/manifest official_text differs from label-100.xlsx")
    text_sha = hashlib.sha256(labels[key].encode("utf-8")).hexdigest()
    if row["official_text_sha256"] != text_sha or row["source_label_xlsx_sha256"] != label_sha:
        raise ValueError("official text/workbook SHA-256 mismatch")
    source_rel = Path(input_row["source_relpath"])
    if source_rel.is_absolute() or ".." in source_rel.parts or source_rel.as_posix() != f"{row['video_id']}/{row['clip_id']}.mp4":
        raise ValueError("source relative path does not match official video_id/clip_id")
    source_path = input_root / source_rel
    if not source_path.is_file() or digest(source_path) != row["source_mp4_sha256"]:
        raise ValueError("current source MP4 missing or SHA-256 differs from manifest")
    if key not in audit_rows:
        raise ValueError("sample key absent from full100 audit")
    audit = audit_rows[key]
    rel = Path(row["output_file"])
    if rel.is_absolute() or ".." in rel.parts:
        raise ValueError(f"unsafe output path: {rel}")
    path = root / rel
    if not path.is_file() or digest(path) != row["output_sha256"] or path.stat().st_size != int(row["output_bytes"]):
        raise ValueError("output file missing or manifest size/hash mismatch")
    with np.load(path, allow_pickle=False) as z:
        required = {
            "schema_version", "config_sha256", "sample_key", "video_id", "clip_id", "official_text",
            "source_video_sha256", "source_label_xlsx_sha256", "alignment_trace_sha256",
            "roberta_model_sha256", "roberta_asset_identity_sha256", "mediapipe_model_sha256", "config_sha256",
            "alignment_mode", "alignment_granularity", "text_audio_correspondence", "text_av_time_mapping_status",
            "text_present", "text_content_valid", "audio_present", "audio_observation_valid", "audio_speech_valid",
            "video_present", "audio_visual_time_valid", "vision_feature_valid", "face_feature_valid", "words",
            "word_char_start", "word_char_end", "word_start_sec", "word_end_sec", "word_time_valid",
            "text_word_feat", "text_word_valid", "word_audio_feat", "word_audio_valid", "word_vision_feat", "word_vision_valid",
            "raw_audio_lld_values", "raw_audio_lld_start_sec", "raw_audio_lld_end_sec", "raw_audio_lld_center_sec", "raw_audio_lld_valid", "raw_audio_lld_feature_names", "opensmile_config_sha256",
            "raw_video_blendshape_values", "raw_video_pts_sec", "raw_video_support_start_sec", "raw_video_support_end_sec", "raw_video_face_feature_valid",
            "raw_video_blendshape_names", "roberta_token_ids", "roberta_token_offsets", "roberta_token_word_index",
            "roberta_word_token_indptr", "roberta_word_token_indices",
            "audio_word_feature_indptr", "audio_word_feature_indices", "vision_word_feature_indptr", "vision_word_feature_indices",
            "audio_presentation_coverage_sec", "video_presentation_coverage_sec", "av_common_coverage_sec",
            "shared_t0_sec", "audio_observation_length", "video_observation_length", "text_sequence_length", "original_effective_duration_sec",
            "alignment_zero_duration_word_count", "alignment_mapping_fail_count",
        }
        missing = required - set(z.files)
        if missing:
            raise ValueError(f"required NPZ fields missing: {sorted(missing)}")
        if str(scalar(z, "sample_key")) != key or str(scalar(z, "schema_version")) != SCHEMA_VERSION:
            raise ValueError("NPZ key/schema mismatch")
        if str(scalar(z, "video_id")) != row["video_id"] or str(scalar(z, "clip_id")) != row["clip_id"]:
            raise ValueError("NPZ video/clip identity mismatch")
        if str(scalar(z, "official_text")) != labels[key]:
            raise ValueError("NPZ official_text differs from workbook")
        if str(scalar(z, "source_video_sha256")) != row["source_mp4_sha256"]:
            raise ValueError("NPZ source MP4 SHA mismatch")
        if str(scalar(z, "source_label_xlsx_sha256")) != label_sha:
            raise ValueError("NPZ label workbook SHA mismatch")
        if str(scalar(z, "roberta_asset_identity_sha256")) != assets.get("roberta_asset_identity_sha256"):
            raise ValueError("NPZ RoBERTa file-set identity differs from pinned asset registry")
        if str(scalar(z, "mediapipe_model_sha256")) != assets.get("mediapipe_model_sha256"):
            raise ValueError("NPZ MediaPipe model identity differs from pinned asset registry")
        if list(z["raw_audio_lld_feature_names"].astype(str)) != assets.get("opensmile_feature_names"):
            raise ValueError("NPZ openSMILE feature order differs from pinned asset registry")
        if str(scalar(z, "alignment_trace_sha256")) != row["alignment_trace_sha256"]:
            raise ValueError("NPZ alignment trace identity differs from manifest")
        if str(scalar(z, "opensmile_config_sha256")) != row["opensmile_config_sha256"]:
            raise ValueError("openSMILE config SHA differs from manifest")
        if str(scalar(z, "roberta_asset_identity_sha256")) != row["roberta_asset_identity_sha256"] or str(scalar(z, "mediapipe_model_sha256")) != row["mediapipe_model_sha256"]:
            raise ValueError("NPZ pinned model asset identity differs from manifest")
        if str(scalar(z, "config_sha256")) != row["config_sha256"]:
            raise ValueError("NPZ frozen config identity differs from manifest")

        words = z["words"].astype(str)
        char_start, char_end = z["word_char_start"], z["word_char_end"]
        expected = expected_words(labels[key])
        if len(words) != len(expected) or not np.array_equal(char_start, [x[0] for x in expected]) or not np.array_equal(char_end, [x[1] for x in expected]) or words.tolist() != [x[2] for x in expected]:
            raise ValueError("NPZ word list/spans do not reconstruct from exact official text")
        L = len(words)
        starts, ends = z["word_start_sec"], z["word_end_sec"]
        wtime = z["word_time_valid"].astype(np.uint8)
        text_feat, text_valid = z["text_word_feat"], z["text_word_valid"].astype(np.uint8)
        audio_feat, audio_valid = z["word_audio_feat"], z["word_audio_valid"].astype(np.uint8)
        vision_feat, vision_valid = z["word_vision_feat"], z["word_vision_valid"].astype(np.uint8)
        if text_feat.shape != (L, 768) or audio_feat.shape != (L, 50) or vision_feat.shape != (L, 104):
            raise ValueError("word-level feature dimensions mismatch")
        if any(mask.shape != (L,) for mask in (wtime, text_valid, audio_valid, vision_valid)):
            raise ValueError("word mask length mismatch")
        if not np.all(text_valid) or not np.isfinite(text_feat).all():
            raise ValueError("official text word representation is not fully finite/valid")
        if np.any(audio_valid & ~wtime) or np.any(vision_valid & ~wtime):
            raise ValueError("word modality valid without word time")
        if np.any(~wtime & (~np.isnan(starts) | ~np.isnan(ends))) or np.any(wtime & (~np.isfinite(starts) | ~np.isfinite(ends) | (ends <= starts))):
            raise ValueError("word time values and masks conflict")
        if np.any(wtime) and (np.any(np.diff(starts[wtime]) < -1e-6) or np.any(np.diff(ends[wtime]) < -1e-6)):
            raise ValueError("word intervals are not monotonic")

        audio = z["raw_audio_lld_values"]
        astart, aend = z["raw_audio_lld_start_sec"], z["raw_audio_lld_end_sec"]
        ac, av = z["raw_audio_lld_center_sec"], z["raw_audio_lld_valid"].astype(np.uint8)
        if audio.ndim != 2 or audio.shape[1] != 25 or not len(audio) or any(x.shape != (len(audio),) for x in (astart, aend, ac, av)):
            raise ValueError("native audio feature schema mismatch")
        if (not np.isfinite(audio).all() or not np.all(av) or not all(np.isfinite(x).all() for x in (astart, aend, ac))
                or np.any(aend <= astart) or np.any(ac < astart) or np.any(ac >= aend)
                or np.any(np.diff(astart) < -1e-6) or np.any(np.diff(ac) < -1e-6)):
            raise ValueError("native audio values/mask/window supports/centers invalid")
        if len(z["raw_audio_lld_feature_names"]) != 25:
            raise ValueError("openSMILE feature-name list length mismatch")
        token_ids, token_offsets, token_word = z["roberta_token_ids"], z["roberta_token_offsets"], z["roberta_token_word_index"]
        token_ptr, token_indices = z["roberta_word_token_indptr"], z["roberta_word_token_indices"]
        if token_offsets.shape != (len(token_ids), 2) or token_word.shape != (len(token_ids),) or token_ptr.shape != (L + 1,) or token_ptr[0] != 0 or token_ptr[-1] != len(token_indices) or np.any(np.diff(token_ptr) < 0):
            raise ValueError("RoBERTa token-to-word mapping arrays have inconsistent CSR shapes")
        seen_token_indices = []
        for wi in range(L):
            indexes = token_indices[token_ptr[wi]:token_ptr[wi + 1]]
            if np.any(indexes < 0) or np.any(indexes >= len(token_ids)) or np.any(token_word[indexes] != wi):
                raise ValueError(f"RoBERTa token CSR assignment differs from token_word_index for word {wi}")
            seen_token_indices.extend(int(x) for x in indexes)
        if len(seen_token_indices) != len(set(seen_token_indices)):
            raise ValueError("a RoBERTa token is assigned to more than one official word")
        official_text = labels[key]
        word_spans = [(int(a), int(b)) for a, b in zip(char_start, char_end)]
        for ti, (a, b) in enumerate(token_offsets.tolist()):
            if a == b:
                if token_word[ti] != -1:
                    raise ValueError("special/empty-offset token must not map to an official word")
                continue
            matches = [wi for wi, (ws, we) in enumerate(word_spans) if ws <= a and b <= we]
            if len(matches) == 1:
                expected_word = matches[0]
            elif not matches and official_text[a:b].isspace():
                expected_word = -1
            else:
                raise ValueError(f"token character span [{a},{b}) cannot be assigned by official word spans")
            if int(token_word[ti]) != expected_word:
                raise ValueError(f"token_word_index does not match official character spans at token {ti}")
        audio_ref, audio_ref_valid, audio_ptr, audio_idx = aggregate_reference(starts, ends, ac, audio, wtime)
        if not np.array_equal(audio_valid, audio_ref_valid) or not np.array_equal(z["audio_word_feature_indptr"], audio_ptr) or not np.array_equal(z["audio_word_feature_indices"], audio_idx) or not np.allclose(audio_feat, audio_ref, rtol=1e-5, atol=1e-6):
            raise ValueError("audio center assignment/CSR/mean/std does not independently recompute")

        video = z["raw_video_blendshape_values"]
        pts, vstart, vend = z["raw_video_pts_sec"], z["raw_video_support_start_sec"], z["raw_video_support_end_sec"]
        face = z["raw_video_face_feature_valid"].astype(np.uint8)
        if video.ndim != 2 or video.shape[1] != 52 or not len(video) or any(a.shape != (len(video),) for a in (pts, vstart, vend, face)):
            raise ValueError("native video feature schema mismatch")
        if not np.isfinite(video).all() or not all(np.isfinite(x).all() for x in (pts, vstart, vend)) or np.any(np.diff(pts) <= 0) or np.any(vend <= vstart) or np.any(pts < vstart - 1e-9) or np.any(pts >= vend + 1e-9) or np.any(video[~face.astype(bool)] != 0):
            raise ValueError("native video values/PTS/support/face mask invalid")
        if abs(vstart[0] - pts[0]) > 1e-9:
            raise ValueError("native video first-frame support does not begin at decoded coverage start")
        if np.any(np.abs(vend[:-1] - vstart[1:]) > 1e-9):
            raise ValueError("native video frame support intervals are not contiguous")
        vision_ref, vision_ref_valid, vision_ptr, vision_idx = aggregate_reference(starts, ends, pts, video, wtime, face)
        if not np.array_equal(vision_valid, vision_ref_valid) or not np.array_equal(z["vision_word_feature_indptr"], vision_ptr) or not np.array_equal(z["vision_word_feature_indices"], vision_idx) or not np.allclose(vision_feat, vision_ref, rtol=1e-5, atol=1e-6):
            raise ValueError("vision center assignment/CSR/mean/std does not independently recompute")

        audio_cov, video_cov, common = z["audio_presentation_coverage_sec"], z["video_presentation_coverage_sec"], z["av_common_coverage_sec"]
        if audio_cov.shape != (2,) or video_cov.shape != (2,) or common.shape != (2,) or audio_cov[1] <= audio_cov[0] or video_cov[1] <= video_cov[0]:
            raise ValueError("presentation coverage fields invalid")
        expected_common = np.asarray([max(audio_cov[0], video_cov[0]), min(audio_cov[1], video_cov[1])])
        if expected_common[1] <= expected_common[0] or not np.allclose(common, expected_common, atol=1e-6, rtol=0):
            raise ValueError("shared A/V coverage does not equal decoded intersection")
        if abs(pts[0] - video_cov[0]) > 1e-9 or abs(vend[-1] - video_cov[1]) > 1e-9:
            raise ValueError("video frame support does not match decoded video coverage")
        shared_t0 = float(scalar(z, "shared_t0_sec"))
        expected_duration = max(float(audio_cov[1]), float(video_cov[1])) - shared_t0
        if expected_duration <= 0 or abs(float(scalar(z, "original_effective_duration_sec")) - expected_duration) > 1e-8:
            raise ValueError("original_effective_duration_sec differs from decoded source coverage")
        if np.any(wtime & ((starts < common[0] - 1e-6) | (ends > common[1] + 1e-6))):
            raise ValueError("word interval outside shared A/V coverage")

        if int(scalar(z, "audio_present")) != 1 or int(scalar(z, "video_present")) != 1 or int(scalar(z, "audio_visual_time_valid")) != 1:
            raise ValueError("audio/video presence or shared timeline flag invalid")
        if int(scalar(z, "text_present")) != 1 or int(scalar(z, "text_content_valid")) != 1 or int(scalar(z, "audio_observation_valid")) != 1:
            raise ValueError("text/audio observation validity flags invalid")
        if int(scalar(z, "face_feature_valid")) != int(face.any()) or int(scalar(z, "vision_feature_valid")) != int(face.any()):
            raise ValueError("face/vision feature status conflicts with actual face mask")
        if int(scalar(z, "audio_speech_valid")) not in {-1, 0, 1}:
            raise ValueError("audio_speech_valid must be -1/0/1")
        if int(scalar(z, "text_sequence_length")) != L or int(scalar(z, "audio_observation_length")) != len(audio) or int(scalar(z, "video_observation_length")) != len(video):
            raise ValueError("stored sequence lengths mismatch actual arrays")
        if any(name in z.files for name in ("sentiment_label", "emotion_label", "target_label", "sentiment", "label")):
            raise ValueError("target labels must not be written into feature artifacts")

        mode, granularity, status = (str(scalar(z, name)) for name in ("alignment_mode", "alignment_granularity", "text_av_time_mapping_status"))
        correspondence = audit.get("text_audio_correspondence", "")
        if correspondence == "confirmed_match":
            if mode != "TRI_MODAL_WORD_VALID" or granularity != "word" or not wtime.any():
                raise ValueError("confirmed match was not represented with its audit-supported word mapping")
            align_path = audit_root / "alignment" / f"{''.join(ch if ch.isalnum() or ch in '-_' else '_' for ch in key).strip('_')}.json"
            alignment = json.loads(align_path.read_text(encoding="utf-8"))
            trace_path = Path(alignment["trace_relative_path"])
            if not trace_path.is_absolute():
                trace_path = audit_root.parents[3] / trace_path
            if not trace_path.is_file() or digest(trace_path) != alignment["trace_sha256"]:
                raise ValueError("audit alignment trace missing or hash mismatch")
            trace = json.loads(trace_path.read_text(encoding="utf-8"))
            trace_words = trace.get("mapped_words", [])
            if len(trace_words) != L:
                raise ValueError("confirmed mapping trace word count differs from official transcript")
            for i, item in enumerate(trace_words):
                valid = bool(item.get("alignment_valid", 0)) and item.get("start_sec") is not None and item.get("end_sec") is not None and float(item["end_sec"]) > float(item["start_sec"])
                if bool(wtime[i]) != valid:
                    raise ValueError(f"word-time validity differs from verified trace at index {i}")
                if valid and (abs(starts[i] - float(item["start_sec"])) > 1e-8 or abs(ends[i] - float(item["end_sec"])) > 1e-8):
                    raise ValueError(f"stored word time differs from verified trace at index {i}")
        else:
            if granularity != "clip" or wtime.any() or status not in {"clip_only", "unavailable"}:
                raise ValueError("unconfirmed/mismatched text mapping fabricated word-level timing")
            if correspondence == "confirmed_mismatch" and mode != "AV_VALID_TEXT_UNALIGNED":
                raise ValueError("confirmed mismatch is not routed to AV_VALID_TEXT_UNALIGNED")
            if correspondence == "no_speech" and mode != "TRI_MODAL_WITH_AUDIO_CONTENT_INVALID":
                raise ValueError("confirmed silence is not retained as text/audio-content-invalid")
            if correspondence not in {"confirmed_mismatch", "no_speech"} and mode != "UNCERTAIN_REVIEW":
                raise ValueError("unconfirmed correspondence is not conservatively clip-only")

        manifest_dims = json.loads(row["feature_dims_json"])
        expected_dims = {"text_word": 768, "audio_native": 25, "video_native": 52, "word_audio_mean_std": 50, "word_vision_mean_std": 104}
        if manifest_dims != expected_dims:
            raise ValueError("manifest feature dimensions differ from stored schema")
        for column, value in (("alignment_mode", mode), ("alignment_granularity", granularity), ("text_av_time_mapping_status", status)):
            if row[column] != value:
                raise ValueError(f"manifest/NPZ mismatch at {column}")
        if int(row["official_word_count"]) != L or int(row["valid_word_time_count"]) != int(wtime.sum()) or int(row["audio_word_valid_count"]) != int(audio_valid.sum()) or int(row["vision_word_valid_count"]) != int(vision_valid.sum()):
            raise ValueError("manifest counts differ from NPZ arrays")
        if row["modalities"] != "text|audio|video":
            raise ValueError("manifest modality presence string must retain all source modalities")
        if abs(float(row["original_effective_duration_sec"]) - float(scalar(z, "original_effective_duration_sec"))) > 1e-9:
            raise ValueError("manifest effective duration differs from NPZ")
        return {
            "sample_key": key, "alignment_mode": mode, "alignment_granularity": granularity,
            "official_word_count": L, "valid_word_time_count": int(wtime.sum()),
            "audio_word_valid_count": int(audio_valid.sum()), "vision_word_valid_count": int(vision_valid.sum()),
            "audio_observation_length": len(audio), "video_observation_length": len(video),
            "face_frame_count": int(face.sum()), "output_sha256": row["output_sha256"],
            "output_bytes": int(row["output_bytes"]),
        }


def main() -> int:
    parser = argparse.ArgumentParser(description="Independent Q1 feature, time, text identity, and center-aggregation validation")
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--input-root", type=Path, required=True)
    parser.add_argument("--audit-root", type=Path, required=True)
    parser.add_argument("--asset-root", type=Path, required=True)
    parser.add_argument("--expected-count", type=int, default=100)
    args = parser.parse_args()
    root, input_root, audit_root = args.output_root.resolve(), args.input_root.resolve(), args.audit_root.resolve()
    manifest_path, input_path, results_path = root / "manifest.csv", root / "input_manifest.csv", root / "results_100.csv"
    if not manifest_path.is_file() or not input_path.is_file() or not results_path.is_file():
        raise SystemExit("FAIL: manifest.csv, input_manifest.csv, or results_100.csv missing")
    manifest, inputs = load_csv(manifest_path), load_csv(input_path)
    missing_columns = REQUIRED_MANIFEST - set(manifest[0] if manifest else [])
    if missing_columns:
        raise SystemExit(f"FAIL: manifest missing required fields: {sorted(missing_columns)}")
    keys, input_keys = [r["sample_key"] for r in manifest], [r["sample_key"] for r in inputs]
    if len(manifest) != args.expected_count or len(set(keys)) != len(keys):
        raise SystemExit(f"FAIL: expected {args.expected_count} unique manifest rows; got {len(manifest)}")
    if len(set(input_keys)) != len(input_keys) or not (set(keys) == set(input_keys) if args.expected_count == 100 else set(keys).issubset(input_keys)):
        raise SystemExit("FAIL: manifest sample-key set does not match the prepared input scope")
    input_by_key = {r["sample_key"]: r for r in inputs}
    config_path = root / "config_snapshot.json"
    assets_path = root / "metadata" / "model_assets.json"
    if not config_path.is_file() or not assets_path.is_file():
        raise SystemExit("FAIL: frozen config or model asset registry missing")
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config_sha = config.get("config_sha256")
    config_payload = {k: v for k, v in config.items() if k != "config_sha256"}
    actual_config_sha = hashlib.sha256(json.dumps(config_payload, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()
    if config_sha != actual_config_sha:
        raise SystemExit("FAIL: config_snapshot.json SHA-256 is internally inconsistent")
    assets = json.loads(assets_path.read_text(encoding="utf-8"))
    asset_root = args.asset_root.resolve()
    model_dir = asset_root / "roberta-base" / assets["roberta_revision"]
    actual_roberta_hashes = {name: digest(model_dir / name) for name in assets["roberta_files_sha256"] if (model_dir / name).is_file()}
    if actual_roberta_hashes != assets["roberta_files_sha256"]:
        raise SystemExit("FAIL: current pinned RoBERTa files missing or hash-mismatched")
    asset_identity = hashlib.sha256(json.dumps(actual_roberta_hashes, sort_keys=True).encode("utf-8")).hexdigest()
    if asset_identity != assets.get("roberta_asset_identity_sha256"):
        raise SystemExit("FAIL: RoBERTa composite asset identity mismatch")
    face_path = asset_root / assets["mediapipe_model_file"]
    if not face_path.is_file() or digest(face_path) != assets.get("mediapipe_model_sha256"):
        raise SystemExit("FAIL: current MediaPipe task model missing or hash-mismatched")
    try:
        from opensmile import FeatureLevel, FeatureSet, Smile
        smile = Smile(feature_set=FeatureSet.eGeMAPSv02, feature_level=FeatureLevel.LowLevelDescriptors,
                      num_workers=1, multiprocessing=False, verbose=False)
        if digest(Path(smile.config_path)) != assets.get("opensmile_config_sha256") or list(smile.feature_names) != assets.get("opensmile_feature_names"):
            raise SystemExit("FAIL: current openSMILE config or feature order differs from frozen asset registry")
    except SystemExit:
        raise
    except Exception as exc:
        raise SystemExit(f"FAIL: cannot verify openSMILE asset identity: {type(exc).__name__}: {exc}")
    result_rows = load_csv(results_path)
    if [(r.get("sample_key"), r.get("output_sha256"), r.get("status")) for r in result_rows] != [(r.get("sample_key"), r.get("output_sha256"), r.get("status")) for r in manifest]:
        raise SystemExit("FAIL: results_100.csv rows differ from manifest.csv")
    labels, label_sha = read_official_labels(input_root / "label-100.xlsx")
    if any(r["source_label_xlsx_sha256"] != label_sha for r in inputs):
        raise SystemExit("FAIL: current label workbook SHA differs from the prepared input manifest")
    audit_rows = {r["sample_key"]: r for r in load_csv(audit_root / "audit_100.csv")}
    for row in manifest:
        if row["status"] != "PASS":
            continue
        if row["config_sha256"] != config_sha:
            raise SystemExit(f"FAIL: manifest config identity differs for {row['sample_key']}")
        if row["roberta_asset_identity_sha256"] != assets.get("roberta_asset_identity_sha256"):
            raise SystemExit(f"FAIL: manifest RoBERTa asset identity differs for {row['sample_key']}")
        if row["mediapipe_model_sha256"] != assets.get("mediapipe_model_sha256"):
            raise SystemExit(f"FAIL: manifest MediaPipe model identity differs for {row['sample_key']}")
        if row["opensmile_config_sha256"] != assets.get("opensmile_config_sha256"):
            raise SystemExit(f"FAIL: manifest openSMILE config identity differs for {row['sample_key']}")
    valid, failures = [], []
    for row in manifest:
        try:
            valid.append(validate_one(root, input_root, row, input_by_key[row["sample_key"]], labels, label_sha, audit_rows, audit_root, assets))
        except Exception as exc:
            failures.append({"sample_key": row.get("sample_key"), "error": f"{type(exc).__name__}: {exc}"})
    modes: dict[str, int] = {}
    for item in valid:
        modes[item["alignment_mode"]] = modes.get(item["alignment_mode"], 0) + 1
    summary = {
        "schema_version": SCHEMA_VERSION, "expected_count": args.expected_count,
        "manifest_count": len(manifest), "valid_count": len(valid), "failure_count": len(failures),
        "status_counts": {status: sum(row["status"] == status for row in manifest) for status in sorted({row["status"] for row in manifest})},
        "alignment_mode_counts_among_valid": modes,
        "feature_file_count": sum(bool(row["output_file"]) for row in manifest),
        "feature_bytes": sum(int(row["output_bytes"] or 0) for row in manifest),
        "label_xlsx_sha256": label_sha, "manifest_sha256": digest(manifest_path),
        "results_sha256": digest(results_path), "failures": failures, "samples": valid,
    }
    reports = root / "reports"
    reports.mkdir(parents=True, exist_ok=True)
    (reports / "automatic_validation.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md = ["# Q1 自动验收报告", "", f"- 期望样本数：{args.expected_count}", f"- manifest 行数：{len(manifest)}", f"- 结构与语义验收通过：{len(valid)}", f"- 失败：{len(failures)}", f"- NPZ 特征总字节数：{summary['feature_bytes']}", f"- 模式计数：`{json.dumps(modes, ensure_ascii=False, sort_keys=True)}`", "", "本报告由独立验证器生成；通过项证明文件身份、结构、时间约束、官方文本映射和聚合复算满足检查规则，不证明未审计文本与真实语音内容一致。", ""]
    if failures:
        md += ["## 失败明细", ""] + [f"- `{x['sample_key']}`：{x['error']}" for x in failures]
    (reports / "automatic_validation.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    if failures:
        print(f"FAIL: {len(failures)} of {len(manifest)} rows failed independent validation")
        for failure in failures[:20]:
            print(f"{failure['sample_key']}: {failure['error']}")
        return 1
    print(f"PASS: {len(valid)} samples; modes={modes}; feature_bytes={summary['feature_bytes']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
