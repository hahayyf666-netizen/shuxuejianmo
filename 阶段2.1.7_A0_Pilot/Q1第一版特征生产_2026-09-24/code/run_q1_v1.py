from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import os
import shutil
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[4]
DELIVERY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT_ROOT = Path(r"D:\Workspace\数学建模\E题数据\附件1-数据集原始多模态样本\MOSEI数据集部分原始视频-100条")
DEFAULT_AUDIT_ROOT = PROJECT_ROOT / "outputs" / "q1" / "diagnostics" / "full100_correspondence_audit"
DEFAULT_ASSET_ROOT = PROJECT_ROOT / "work" / "model_assets_a0"
STAGE1_CSV = PROJECT_ROOT / "outputs" / "q1" / "pilot_selection_metadata.csv"
ROBERTA_REVISION = "e2da8e2f811d1448a5b465c236feacd80ffbac7b"
ROBERTA_SHA256 = "5bde1d28afb363d0103324efeb5afc8b2b397fe5e04beabb9b1ef355255ade81"
FACE_MODEL_SHA256 = "64184e229b263107bc2b804c6625db1341ff2bb731874b0bcc2fe6544e0bc9ff"
ROBERTA_FILE_SHA256 = {
    "config.json": "ef0185e2aae6e06c5f105a285006952c340e20c7dbf43c86ec82601b13fc45e9",
    "merges.txt": "1ce1664773c50f3e0cc8842619a93edc4624525b728b188a9e0be33b7726adc5",
    "model.safetensors": "5bde1d28afb363d0103324efeb5afc8b2b397fe5e04beabb9b1ef355255ade81",
    "tokenizer.json": "847bbeab6174d66a88898f729d52fa8d355fafe1bea101cf960dd404581df70e",
    "tokenizer_config.json": "994f46754c5bf4014f1aa92d34b1374319c3a6b3f702105cd5b742beaecd18ce",
    "vocab.json": "9e7f63c2d15d666b52e21d250d2e513b87c9b713cfa6987a82ed89e5e6e50655",
}
SCHEMA_VERSION = "q1-feature-v1.0"


class PipelineStop(RuntimeError):
    pass


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    temp.replace(path)


def safe_stem(sample_key: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in sample_key).strip("_")


def norm_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def load_registry(input_root: Path, audit_root: Path) -> list[dict[str, Any]]:
    audit_csv = audit_root / "audit_100.csv"
    summary_path = audit_root / "audit_summary.json"
    if not audit_csv.is_file() or not summary_path.is_file() or not STAGE1_CSV.is_file():
        raise PipelineStop("required full100 audit or Stage1 selection files are missing")
    audit_rows = read_csv(audit_csv)
    stage1_rows = read_csv(STAGE1_CSV)
    if len(audit_rows) != 100 or len(stage1_rows) != 100:
        raise PipelineStop(f"expected exactly 100 audit and Stage1 rows; got {len(audit_rows)} and {len(stage1_rows)}")
    audit_by_key = {row["sample_key"]: row for row in audit_rows}
    stage1_by_key = {row["sample_key"]: row for row in stage1_rows}
    if len(audit_by_key) != 100 or len(stage1_by_key) != 100 or set(audit_by_key) != set(stage1_by_key):
        raise PipelineStop("audit and Stage1 sample keys are not one-to-one")

    label_path = input_root / "label-100.xlsx"
    if not label_path.is_file():
        raise PipelineStop(f"official label workbook missing: {label_path}")
    label_sha = sha256(label_path)
    audit_summary = json.loads(summary_path.read_text(encoding="utf-8"))
    if label_sha != audit_summary.get("official_label_sha256"):
        raise PipelineStop("label-100.xlsx SHA-256 differs from the full100 audit source")

    workbook = load_workbook(label_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(iterator)]
    needed = {"video_id", "clip_id", "text"}
    if not needed.issubset(header):
        workbook.close()
        raise PipelineStop(f"official label table lacks columns: {sorted(needed - set(header))}")
    indexes = {name: header.index(name) for name in needed}
    labels: dict[tuple[str, str], str] = {}
    for row in iterator:
        key = (norm_id(row[indexes["video_id"]]), norm_id(row[indexes["clip_id"]]))
        if key in labels:
            workbook.close()
            raise PipelineStop(f"duplicate official label row for {key}")
        labels[key] = row[indexes["text"]]
    workbook.close()

    summary_dir = audit_root / "alignment"
    media_dir = audit_root / "media"
    registry: list[dict[str, Any]] = []
    for key in sorted(audit_by_key):
        row = audit_by_key[key]
        stage1 = stage1_by_key[key]
        video_id, clip_id = norm_id(row["video_id"]), norm_id(row["clip_id"])
        expected_key = f"{video_id}$_${clip_id}"
        if key != expected_key:
            raise PipelineStop(f"sample key format mismatch for {key}; expected {expected_key}")
        text = labels.get((video_id, clip_id))
        if not isinstance(text, str) or not text.strip() or text != row["official_text"]:
            raise PipelineStop(f"official_text differs from label-100.xlsx for {key}")
        relative_video = f"{video_id}/{clip_id}.mp4"
        video_path = input_root / video_id / f"{clip_id}.mp4"
        if not video_path.is_file():
            raise PipelineStop(f"source MP4 missing for {key}: {relative_video}")
        video_sha = sha256(video_path)
        if video_sha != row["source_mp4_sha256"]:
            raise PipelineStop(f"source MP4 SHA-256 changed for {key}")

        stub = safe_stem(key)
        media_path = media_dir / f"{stub}.json"
        alignment_path = summary_dir / f"{stub}.json"
        if not media_path.is_file() or not alignment_path.is_file():
            raise PipelineStop(f"audit evidence missing for {key}")
        media = json.loads(media_path.read_text(encoding="utf-8"))
        alignment = json.loads(alignment_path.read_text(encoding="utf-8"))
        if media.get("source_mp4_sha256") != video_sha or alignment.get("source_mp4_sha256") != video_sha:
            raise PipelineStop(f"audit evidence is for a different source MP4: {key}")
        if alignment.get("official_text_sha256") != hashlib.sha256(text.encode("utf-8")).hexdigest():
            raise PipelineStop(f"alignment evidence is for different official text: {key}")
        if alignment.get("alignment_diagnostic_status") != "completed":
            raise PipelineStop(f"alignment diagnostic is not complete for {key}")
        trace_rel = alignment.get("trace_relative_path")
        if not trace_rel:
            raise PipelineStop(f"alignment trace path is missing for {key}")
        trace_path = PROJECT_ROOT / Path(trace_rel)
        if not trace_path.is_file() or sha256(trace_path) != alignment.get("trace_sha256"):
            raise PipelineStop(f"alignment trace is missing or hash-mismatched for {key}")
        trace = json.loads(trace_path.read_text(encoding="utf-8"))
        if trace.get("sample_key") != key or trace.get("official_transcript") != text:
            raise PipelineStop(f"alignment trace identity/text mismatch for {key}")
        if int(row["audio_visual_time_valid"]) != 1:
            raise PipelineStop(f"full100 audit does not certify a shared A/V timeline for {key}")
        audio_start, video_start = media.get("audio_start_sec"), media.get("video_start_sec")
        audio_end, video_end = media.get("audio_end_sec"), media.get("video_end_sec")
        if any(value is None for value in (audio_start, video_start, audio_end, video_end)):
            raise PipelineStop(f"A/V presentation coverage is incomplete for {key}")
        if abs(float(audio_start) - float(video_start)) > 1e-6 or min(float(audio_end), float(video_end)) <= max(float(audio_start), float(video_start)):
            raise PipelineStop(f"A/V shared presentation timeline is not reconcilable for {key}")
        registry.append({
            "sample_key": key,
            "video_id": video_id,
            "clip_id": clip_id,
            "source_relpath": relative_video,
            "source_mp4_sha256": video_sha,
            "source_mp4_bytes": video_path.stat().st_size,
            "official_text": text,
            "official_text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
            "source_label_xlsx_sha256": label_sha,
            "duration_proxy_sec": float(stage1["ffprobe_format_duration_sec"]),
            "word_count_proxy": int(stage1["word_count_proxy"]),
            "audit_row": row,
            "media_evidence": media,
            "alignment_evidence": alignment,
            "alignment_trace": trace,
            "alignment_trace_sha256": alignment["trace_sha256"],
            "shared_t0_sec": float(audio_start),
        })
    return registry


def check_assets(asset_root: Path) -> dict[str, Any]:
    from opensmile import FeatureLevel, FeatureSet, Smile
    model_dir = asset_root / "roberta-base" / ROBERTA_REVISION
    face_path = asset_root / "mediapipe" / "face_landmarker.task"
    observed_hashes = {}
    for filename, expected_sha in ROBERTA_FILE_SHA256.items():
        path = model_dir / filename
        if not path.is_file() or sha256(path) != expected_sha:
            raise PipelineStop(f"pinned RoBERTa asset missing or SHA-256 mismatch: {filename}")
        observed_hashes[filename] = sha256(path)
    if not face_path.is_file() or sha256(face_path) != FACE_MODEL_SHA256:
        raise PipelineStop("pinned MediaPipe task model missing or SHA-256 mismatch")
    smile = Smile(feature_set=FeatureSet.eGeMAPSv02, feature_level=FeatureLevel.LowLevelDescriptors,
                  num_workers=1, multiprocessing=False, verbose=False)
    opensmile_config_path = Path(smile.config_path)
    opensmile_config_sha = sha256(opensmile_config_path)
    opensmile_feature_names = list(smile.feature_names)
    if len(opensmile_feature_names) != 25:
        raise PipelineStop("pinned openSMILE eGeMAPSv02 LLD feature-name schema is not 25D")
    if sys.version_info[:3] != (3, 12, 10):
        raise PipelineStop(f"Python version differs from qualified environment: {sys.version_info[:3]}")
    expected_versions = {
        "torch": "2.8.0", "torchaudio": "2.8.0", "stable-ts": "2.19.1",
        "openai-whisper": "20250625", "opensmile": "2.6.0", "mediapipe": "0.10.35",
        "transformers": "4.57.6", "av": "18.0.0", "pandas": "3.0.3", "openpyxl": "3.1.5",
    }
    observed_versions = {name: importlib.metadata.version(name) for name in expected_versions}
    if observed_versions != expected_versions:
        raise PipelineStop(f"package versions differ from qualified environment: {observed_versions}")
    identity_sha = hashlib.sha256(json.dumps(observed_hashes, sort_keys=True).encode("utf-8")).hexdigest()
    return {
        "roberta_model_id": "FacebookAI/roberta-base",
        "roberta_revision": ROBERTA_REVISION,
        "roberta_files_sha256": observed_hashes,
        "roberta_model_file": "roberta-base/<revision>/model.safetensors",
        "roberta_model_sha256": observed_hashes["model.safetensors"],
        "roberta_asset_identity_sha256": identity_sha,
        "mediapipe_model_file": "mediapipe/face_landmarker.task",
        "mediapipe_model_sha256": sha256(face_path),
        "opensmile_config_file": str(opensmile_config_path.name),
        "opensmile_config_sha256": opensmile_config_sha,
        "opensmile_feature_names": opensmile_feature_names,
        "stable_ts_alignment_reused_from_audit": True,
        "stable_ts_version": importlib.metadata.version("stable-ts"),
        "openai_whisper_version": importlib.metadata.version("openai-whisper"),
        "qualified_package_versions": observed_versions,
        "whisper_checkpoint_loaded_in_feature_run": False,
    }


def make_config(asset_info: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "audio": {"decode": "PyAV decoded source audio with presentation PTS", "sample_rate_hz": 16000, "channels": 1, "feature_set": "eGeMAPSv02", "feature_level": "LowLevelDescriptors", "feature_dim": 25, "opensmile_config_sha256": asset_info["opensmile_config_sha256"], "feature_names": asset_info["opensmile_feature_names"]},
        "video": {"decode": "PyAV decoded frame presentation PTS", "feature_extractor": "MediaPipe FaceLandmarker", "mode": "VIDEO", "num_faces": 1, "blendshape_dim": 52, "create_fresh_instance_per_sample": True},
        "text": {"source": "label-100.xlsx / text", "encoder": "FacebookAI/roberta-base", "revision": ROBERTA_REVISION, "asset_identity_sha256": asset_info["roberta_asset_identity_sha256"], "pool": "mean of non-special subword last-hidden-state vectors within exact official-word character span", "hidden_dim": 768, "truncate": False},
        "alignment": {"timebase": "shared presentation timeline", "word_interval": "half-open [start,end)", "reuse": "existing stable-ts 2.19.1 diagnostic only when source/text hashes and confirmed-match evidence agree", "aggregation_query": "feature center >= word start and < word end", "std_ddof": 0},
        "missingness": {"unavailable_word_times": "NaN with word_time_valid=0", "no_face": "preserve frame with zero placeholder and face_feature_valid=0", "silent_audio": "preserve extracted LLD values; audio_speech_valid=0 only for confirmed silence", "unasserted_text_audio_link": "clip-level text association; no word timestamps"},
        "runtime": {"device": "CPU", "model_inference_batch_size": 1, "media_processing": "serial by sample_key", "stable_ts_transcribe": False, "new_alignment_run": False},
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = list(rows[0]) if rows else []
    temp = path.with_suffix(path.suffix + ".tmp")
    with temp.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temp.replace(path)


def output_manifest_row(row: dict[str, Any], log: dict[str, Any], output_root: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "sample_key": row["sample_key"], "video_id": row["video_id"], "clip_id": row["clip_id"],
        "official_text": row["official_text"], "official_text_sha256": row["official_text_sha256"],
        "source_relpath": row["source_relpath"], "source_mp4_sha256": row["source_mp4_sha256"],
        "source_label_xlsx_sha256": row["source_label_xlsx_sha256"],
        "stage1_duration_proxy_sec": row["duration_proxy_sec"],
        "alignment_trace_sha256": row["alignment_trace_sha256"],
        "status": log.get("status", "MISSING"), "output_file": "", "output_sha256": "", "output_bytes": 0,
    }
    if log.get("status") != "PASS":
        result.update({
            "modalities": "text|audio|video", "original_effective_duration_sec": "",
            "feature_dims_json": json.dumps({"text_word": 768, "audio_native": 25, "video_native": 52, "word_audio_mean_std": 50, "word_vision_mean_std": 104}, sort_keys=True),
            "text_sequence_length": "", "audio_observation_length": "", "video_observation_length": "",
            "alignment_mode": "", "alignment_granularity": "", "text_audio_correspondence": "",
            "text_av_time_mapping_status": "", "text_present": 1, "text_content_valid": 1,
            "audio_present": 1, "audio_observation_valid": "", "audio_speech_valid": -1,
            "video_present": 1, "audio_visual_time_valid": "", "vision_feature_valid": "",
            "face_feature_valid": "", "official_word_count": "", "valid_word_time_count": "",
            "text_word_valid_count": "", "audio_word_valid_count": "", "vision_word_valid_count": "",
            "audio_coverage_start_sec": "", "audio_coverage_end_sec": "",
            "video_coverage_start_sec": "", "video_coverage_end_sec": "",
            "av_common_coverage_start_sec": "", "av_common_coverage_end_sec": "",
            "opensmile_config_sha256": "", "roberta_asset_identity_sha256": "",
            "config_sha256": config_digest(output_root), "mediapipe_model_sha256": "",
            "process_rss_bytes": "", "runtime_sec": log.get("runtime_sec", ""),
            "error": log.get("error", ""),
        })
        return result
    rel = Path(log["output_file"])
    with np.load(output_root / rel, allow_pickle=False) as z:
        scalar = lambda name: z[name].item()
        dims = {
            "text_word": int(z["text_word_feat"].shape[1]),
            "audio_native": int(z["raw_audio_lld_values"].shape[1]),
            "video_native": int(z["raw_video_blendshape_values"].shape[1]),
            "word_audio_mean_std": int(z["word_audio_feat"].shape[1]),
            "word_vision_mean_std": int(z["word_vision_feat"].shape[1]),
        }
        result.update({
            "output_file": str(rel).replace("\\", "/"), "output_sha256": log["output_sha256"],
            "output_bytes": int(log["output_bytes"]), "modalities": "text|audio|video",
            "original_effective_duration_sec": float(scalar("original_effective_duration_sec")),
            "feature_dims_json": json.dumps(dims, sort_keys=True),
            "text_sequence_length": int(scalar("text_sequence_length")),
            "audio_observation_length": int(scalar("audio_observation_length")),
            "video_observation_length": int(scalar("video_observation_length")),
            "alignment_mode": str(scalar("alignment_mode")), "alignment_granularity": str(scalar("alignment_granularity")),
            "text_audio_correspondence": str(scalar("text_audio_correspondence")),
            "text_av_time_mapping_status": str(scalar("text_av_time_mapping_status")),
            "text_present": int(scalar("text_present")), "text_content_valid": int(scalar("text_content_valid")),
            "audio_present": int(scalar("audio_present")), "audio_observation_valid": int(scalar("audio_observation_valid")),
            "audio_speech_valid": int(scalar("audio_speech_valid")), "video_present": int(scalar("video_present")),
            "audio_visual_time_valid": int(scalar("audio_visual_time_valid")),
            "vision_feature_valid": int(scalar("vision_feature_valid")), "face_feature_valid": int(scalar("face_feature_valid")),
            "official_word_count": int(len(z["words"])), "valid_word_time_count": int(z["word_time_valid"].sum()),
            "text_word_valid_count": int(z["text_word_valid"].sum()),
            "audio_word_valid_count": int(z["word_audio_valid"].sum()), "vision_word_valid_count": int(z["word_vision_valid"].sum()),
            "audio_coverage_start_sec": float(z["audio_presentation_coverage_sec"][0]),
            "audio_coverage_end_sec": float(z["audio_presentation_coverage_sec"][1]),
            "video_coverage_start_sec": float(z["video_presentation_coverage_sec"][0]),
            "video_coverage_end_sec": float(z["video_presentation_coverage_sec"][1]),
            "av_common_coverage_start_sec": float(z["av_common_coverage_sec"][0]),
            "av_common_coverage_end_sec": float(z["av_common_coverage_sec"][1]),
            "opensmile_config_sha256": str(scalar("opensmile_config_sha256")),
            "roberta_asset_identity_sha256": str(scalar("roberta_asset_identity_sha256")),
            "config_sha256": str(scalar("config_sha256")),
            "mediapipe_model_sha256": str(scalar("mediapipe_model_sha256")),
            "process_rss_bytes": log.get("process_rss_bytes", ""), "runtime_sec": log.get("runtime_sec", ""), "error": "",
        })
    return result


def rebuild_output_tables(output_root: Path, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    manifest_rows = []
    for row in rows:
        stub = safe_stem(row["sample_key"])
        log_path = output_root / "logs" / f"{stub}.media.json"
        log = json.loads(log_path.read_text(encoding="utf-8")) if log_path.is_file() else {"status": "MISSING", "error": "media worker log missing"}
        manifest_rows.append(output_manifest_row(row, log, output_root))
    fields = list(manifest_rows[0]) if manifest_rows else []
    write_csv(output_root / "manifest.csv", manifest_rows, fields)
    write_csv(output_root / "results_100.csv", manifest_rows, fields)
    return manifest_rows


def baseline_hashes() -> dict[str, str]:
    protected = [
        PROJECT_ROOT / "outputs" / "q1" / "pilot_samples.json",
        PROJECT_ROOT / "outputs" / "q1" / "pilot",
    ]
    result: dict[str, str] = {}
    for item in protected:
        if item.is_file():
            result[str(item.relative_to(PROJECT_ROOT))] = sha256(item)
        elif item.is_dir():
            for file_path in sorted(item.rglob("*")):
                if file_path.is_file():
                    result[str(file_path.relative_to(PROJECT_ROOT))] = sha256(file_path)
    return result


def prepare(args: argparse.Namespace) -> list[dict[str, Any]]:
    output_root = Path(args.output_root).resolve()
    for name in ("features", "logs", "metadata", "reports", ".cache/text"):
        (output_root / name).mkdir(parents=True, exist_ok=True)
    registry = load_registry(Path(args.input_root), Path(args.audit_root))
    if len(registry) != 100:
        raise PipelineStop("input registry did not resolve to 100 samples")
    if args.sample_key:
        requested = args.sample_key
        if len(requested) != len(set(requested)) or not set(requested).issubset({row["sample_key"] for row in registry}):
            raise PipelineStop("requested sample keys contain duplicates or values outside the 100-row universe")
        selected = set(requested)
    else:
        selected = {row["sample_key"] for row in registry}
    selected_registry = [row for row in registry if row["sample_key"] in selected]
    asset_info = check_assets(Path(args.asset_root))
    config = make_config(asset_info)
    config_sha = hashlib.sha256(json.dumps(config, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()
    config["config_sha256"] = config_sha
    snapshot_path = output_root / "config_snapshot.json"
    if args.resume and snapshot_path.is_file():
        frozen = json.loads(snapshot_path.read_text(encoding="utf-8"))
        if frozen != config:
            raise PipelineStop("resume requested with changed configuration; use a new output directory")
    else:
        write_json(snapshot_path, config)
    registry_path = output_root / "metadata" / "registry.json"
    if args.resume and registry_path.is_file():
        prior_registry = json.loads(registry_path.read_text(encoding="utf-8"))
        if [(r["sample_key"], r["source_mp4_sha256"], r["official_text_sha256"]) for r in prior_registry] != [(r["sample_key"], r["source_mp4_sha256"], r["official_text_sha256"]) for r in registry]:
            raise PipelineStop("resume requested with changed input registry; use a new output directory")
    else:
        write_json(registry_path, registry)
    selected_path = output_root / "metadata" / "selected_sample_keys.json"
    if args.resume and selected_path.is_file():
        prior_selected = json.loads(selected_path.read_text(encoding="utf-8"))
        current_selected = [row["sample_key"] for row in selected_registry]
        if prior_selected != current_selected:
            raise PipelineStop("resume requested with changed selected sample set; use a new output directory")
    else:
        write_json(selected_path, [row["sample_key"] for row in selected_registry])
    model_assets_path = output_root / "metadata" / "model_assets.json"
    if args.resume and model_assets_path.is_file() and json.loads(model_assets_path.read_text(encoding="utf-8")) != asset_info:
        raise PipelineStop("resume requested with changed model/feature assets; use a new output directory")
    write_json(model_assets_path, asset_info)
    protected_path = output_root / "metadata" / "protected_a0_baseline.json"
    current_baseline = baseline_hashes()
    if args.resume and protected_path.is_file():
        if json.loads(protected_path.read_text(encoding="utf-8")) != current_baseline:
            raise PipelineStop("protected A0 pilot bytes changed since first prepare")
    else:
        write_json(protected_path, current_baseline)
    evidence_source = DELIVERY_ROOT / "metadata" / "official_answer_evidence"
    evidence_output = output_root / "metadata" / "official_answer_evidence"
    evidence_files = []
    if evidence_source.is_dir():
        for path in sorted(evidence_source.glob("*.png")):
            target = evidence_output / path.name
            if path.resolve() != target.resolve():
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, target)
            evidence_files.append({"file": f"metadata/official_answer_evidence/{path.name}", "sha256": sha256(target), "bytes": target.stat().st_size})
    if len(evidence_files) != 3:
        raise PipelineStop("the three user-provided official answer screenshots are not available in the delivery evidence folder")
    write_json(output_root / "metadata" / "official_answer_evidence.json", {
        "source_type": "user-provided screenshots of the official answer",
        "files": evidence_files,
        "decision_rules": [
            "retain all 100 original samples; do not rewrite official text or replace it with ASR",
            "use the original MP4 audio and video; extract real acoustic features for silent/low-energy audio",
            "retain video frames when no face is detected and mark face feature validity separately",
            "for confirmed text/audio mismatch, keep official text and source A/V while withholding fabricated word timestamps",
        ],
    })
    input_rows = [{key: row[key] for key in (
        "sample_key", "video_id", "clip_id", "source_relpath", "source_mp4_sha256", "source_mp4_bytes",
        "official_text", "official_text_sha256", "source_label_xlsx_sha256", "duration_proxy_sec", "word_count_proxy",
        "shared_t0_sec", "alignment_trace_sha256",
    )} for row in registry]
    write_csv(output_root / "input_manifest.csv", input_rows)
    write_json(output_root / "metadata" / "evidence_manifest.json", {
        "sample_count": len(registry),
        "source_root_argument": str(Path(args.input_root)),
        "audit_root_argument": str(Path(args.audit_root)),
        "official_label_workbook": {"file": "label-100.xlsx", "sha256": registry[0]["source_label_xlsx_sha256"]},
        "full100_audit": {
            "audit_100.csv": sha256(Path(args.audit_root) / "audit_100.csv"),
            "audit_rules.md": sha256(Path(args.audit_root) / "audit_rules.md"),
            "audit_summary.json": sha256(Path(args.audit_root) / "audit_summary.json"),
        },
        "stage1_selection_metadata_sha256": sha256(STAGE1_CSV),
        "official_answer_screenshots": evidence_files,
        "source_mp4s": [{"sample_key": row["sample_key"], "relative_path": row["source_relpath"], "sha256": row["source_mp4_sha256"], "bytes": row["source_mp4_bytes"]} for row in registry],
        "alignment_traces": [{"sample_key": row["sample_key"], "trace_sha256": row["alignment_trace_sha256"], "trace_path": row["alignment_evidence"]["trace_relative_path"]} for row in registry],
        "protected_a0_baseline_sha256": current_baseline,
    })

    import psutil
    import torch
    try:
        import torchaudio
        torchaudio_version = torchaudio.__version__
    except Exception:
        torchaudio_version = None
    cuda_available = bool(torch.cuda.is_available())
    cuda_devices = [torch.cuda.get_device_name(i) for i in range(torch.cuda.device_count())] if cuda_available else []
    memory = psutil.virtual_memory()
    resource = {
        "captured_utc": datetime.now(timezone.utc).isoformat(),
        "ram_total_bytes": int(memory.total),
        "ram_available_bytes": int(memory.available),
        "cpu_count": os.cpu_count(),
        "cuda_available": cuda_available,
        "cuda_devices": cuda_devices,
        "recommended_available_ram_bytes": 4 * 1024**3,
        "policy": "warn and use isolated serial stages below the 4 GiB reserve recommendation; stop before a sample if memory falls below 1 GiB",
        "full_dataset_loaded_into_memory": False,
        "selected_sample_count": len(selected_registry),
    }
    write_json(output_root / "resource_preflight.json", resource)
    env = {
        "sys_executable": sys.executable,
        "python_version": sys.version,
        "platform": sys.platform,
        "cpu_count": os.cpu_count(),
        "ram_total_bytes": int(memory.total),
        "ram_available_bytes_at_preflight": int(memory.available),
        "device": "cpu",
        "torch_version": torch.__version__,
        "torchaudio_version": torchaudio_version,
        "cuda_available": cuda_available,
        "cuda_devices": cuda_devices,
        "input_root_argument": str(Path(args.input_root)),
        "audit_root_relative": str(Path(args.audit_root).resolve().relative_to(PROJECT_ROOT)),
        "schema_version": SCHEMA_VERSION,
        "selected_count": len(selected_registry),
        "label_xlsx_sha256": registry[0]["source_label_xlsx_sha256"],
        "audit_csv_sha256": sha256(Path(args.audit_root) / "audit_100.csv"),
        "audit_rules_sha256": sha256(Path(args.audit_root) / "audit_rules.md"),
        **{f"package_{name.replace('-', '_')}": importlib.metadata.version(name) for name in ("numpy", "torch", "stable-ts", "openai-whisper", "opensmile", "mediapipe", "transformers", "av", "pandas", "openpyxl")},
    }
    write_json(output_root / "environment.json", env)
    freeze = subprocess.run([sys.executable, "-m", "pip", "freeze"], capture_output=True, text=True, check=False)
    (output_root / "requirements-lock.txt").write_text(freeze.stdout, encoding="utf-8")
    if freeze.returncode != 0:
        raise PipelineStop(f"pip freeze failed with exit code {freeze.returncode}: {freeze.stderr[-1000:]}")
    return selected_registry


def import_core():
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import q1_feature_core as core
    return core


def selected_rows(output_root: Path, requested_keys: list[str] | None = None) -> list[dict[str, Any]]:
    registry = json.loads((output_root / "metadata" / "registry.json").read_text(encoding="utf-8"))
    keys = set(json.loads((output_root / "metadata" / "selected_sample_keys.json").read_text(encoding="utf-8")))
    if requested_keys:
        requested = set(requested_keys)
        if len(requested) != len(requested_keys) or not requested.issubset(keys):
            raise PipelineStop("requested sample keys are duplicated or outside the prepared output scope")
        keys &= requested
    return [row for row in registry if row["sample_key"] in keys]


def verify_live_label(input_root: Path, rows: list[dict[str, Any]]) -> None:
    label_path = input_root / "label-100.xlsx"
    if not label_path.is_file() or sha256(label_path) != rows[0]["source_label_xlsx_sha256"]:
        raise PipelineStop("official label workbook changed after preflight")


def load_model_paths(asset_root: Path):
    from transformers import AutoModel, AutoTokenizer
    model_path = asset_root / "roberta-base" / ROBERTA_REVISION
    tokenizer = AutoTokenizer.from_pretrained(str(model_path), local_files_only=True, use_fast=True)
    model = AutoModel.from_pretrained(str(model_path), local_files_only=True, use_safetensors=True).to("cpu").eval()
    if not tokenizer.is_fast or int(model.config.hidden_size) != 768:
        raise PipelineStop("RoBERTa tokenizer/model API differs from the frozen contract")
    return tokenizer, model


def config_digest(output_root: Path) -> str:
    config_path = output_root / "config_snapshot.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    return str(config["config_sha256"])


def validate_text_cache(path: Path, row: dict[str, Any], core, model_sha256: str, model_identity_sha256: str, expected_config_sha256: str) -> dict[str, np.ndarray]:
    if not path.is_file():
        raise PipelineStop(f"text cache missing for {row['sample_key']}")
    with np.load(path, allow_pickle=False) as cached:
        arrays = {name: cached[name].copy() for name in cached.files}
    required = {
        "cache_schema_version", "cache_sample_key", "cache_official_text_sha256", "cache_roberta_revision",
        "cache_roberta_sha256", "cache_roberta_asset_identity_sha256", "cache_config_sha256", "words", "word_char_start", "word_char_end",
        "text_word_feat", "text_word_valid", "roberta_token_ids", "roberta_token_offsets",
        "roberta_token_word_index", "roberta_word_token_indptr", "roberta_word_token_indices",
    }
    missing = required - set(arrays)
    if missing:
        raise PipelineStop(f"text cache fields missing for {row['sample_key']}: {sorted(missing)}")
    meta = {name: str(arrays[name].item()) for name in (
        "cache_schema_version", "cache_sample_key", "cache_official_text_sha256", "cache_roberta_revision",
        "cache_roberta_sha256", "cache_roberta_asset_identity_sha256", "cache_config_sha256",
    )}
    expected = {
        "cache_schema_version": SCHEMA_VERSION,
        "cache_sample_key": row["sample_key"],
        "cache_official_text_sha256": row["official_text_sha256"],
        "cache_roberta_revision": ROBERTA_REVISION,
        "cache_roberta_sha256": model_sha256,
        "cache_roberta_asset_identity_sha256": model_identity_sha256,
        "cache_config_sha256": expected_config_sha256,
    }
    if meta != expected:
        raise PipelineStop(f"text cache identity/config mismatch for {row['sample_key']}: {meta}")
    source_words = core.source_words(row["official_text"])
    expected_words = np.asarray([word["text"] for word in source_words], dtype=np.str_)
    expected_starts = np.asarray([word["char_start"] for word in source_words], dtype=np.int32)
    expected_ends = np.asarray([word["char_end"] for word in source_words], dtype=np.int32)
    if not (np.array_equal(arrays["words"], expected_words)
            and np.array_equal(arrays["word_char_start"], expected_starts)
            and np.array_equal(arrays["word_char_end"], expected_ends)):
        raise PipelineStop(f"text cache official-word sequence/span mismatch for {row['sample_key']}")
    L = len(expected_words)
    if arrays["text_word_feat"].shape != (L, 768) or arrays["text_word_valid"].shape != (L,) or not np.isfinite(arrays["text_word_feat"]).all():
        raise PipelineStop(f"text cache feature shape/value mismatch for {row['sample_key']}")
    if not np.all(arrays["text_word_valid"] == 1):
        raise PipelineStop(f"text cache contains invalid official word features for {row['sample_key']}")
    return arrays


def text_stage(args: argparse.Namespace) -> int:
    output_root = Path(args.output_root).resolve()
    rows = selected_rows(output_root, args.sample_key)
    verify_live_label(Path(args.input_root), rows)
    core = import_core()
    tokenizer, model = load_model_paths(Path(args.asset_root))
    import torch
    torch.set_num_threads(1)
    asset_info = check_assets(Path(args.asset_root))
    saved_assets = json.loads((output_root / "metadata" / "model_assets.json").read_text(encoding="utf-8"))
    if asset_info != saved_assets:
        raise PipelineStop("model/tokenizer assets changed after prepare")
    model_sha = asset_info["roberta_model_sha256"]
    model_identity_sha = asset_info["roberta_asset_identity_sha256"]
    cfg_sha = config_digest(output_root)
    results = []
    hard_stop = False
    for index, row in enumerate(rows, 1):
        key, stub = row["sample_key"], safe_stem(row["sample_key"])
        log_path = output_root / "logs" / f"{stub}.text.json"
        temp_path = output_root / ".cache" / "text" / f"{stub}.npz"
        final_path = output_root / "features" / f"{stub}.npz"
        if args.resume and temp_path.is_file():
            try:
                validate_text_cache(temp_path, row, core, model_sha, model_identity_sha, cfg_sha)
                previous = json.loads(log_path.read_text(encoding="utf-8")) if log_path.is_file() else {}
                if previous.get("text_cache_sha256") != sha256(temp_path):
                    raise PipelineStop(f"text cache file hash differs from its sidecar for {key}")
                resumed = {"sample_key": key, "stage": "text", "status": "RESUMED_TEXT_CACHE", "text_cache_sha256": sha256(temp_path), "official_text_sha256": row["official_text_sha256"], "roberta_model_sha256": model_sha, "roberta_asset_identity_sha256": model_identity_sha, "config_sha256": cfg_sha}
                write_json(log_path, resumed)
                results.append(resumed)
                continue
            except Exception as exc:
                temp_path.unlink(missing_ok=True)
                print(f"TEXT_CACHE_REJECTED {key}: {type(exc).__name__}: {exc}; recomputing", flush=True)
        started = time.perf_counter()
        record: dict[str, Any] = {"sample_key": key, "stage": "text", "status": "RUNNING", "official_text_sha256": row["official_text_sha256"]}
        try:
            words = core.source_words(row["official_text"])
            info = core.extract_text_features(row["official_text"], words, tokenizer, model, torch.device("cpu"))
            if info["tokenizer_failures"]:
                record["status"] = "SAMPLE_FAIL"
                record["error"] = "tokenizer_character_span_mapping_failure"
            else:
                temp_path.parent.mkdir(parents=True, exist_ok=True)
                temp = temp_path.with_suffix(".npz.tmp")
                with temp.open("wb") as stream:
                    np.savez_compressed(
                        stream,
                        cache_schema_version=np.asarray(SCHEMA_VERSION),
                        cache_sample_key=np.asarray(key),
                        cache_official_text_sha256=np.asarray(row["official_text_sha256"]),
                        cache_roberta_revision=np.asarray(ROBERTA_REVISION),
                        cache_roberta_sha256=np.asarray(model_sha),
                        cache_roberta_asset_identity_sha256=np.asarray(model_identity_sha),
                        cache_config_sha256=np.asarray(cfg_sha),
                        words=np.asarray([word["text"] for word in words], dtype=np.str_),
                        word_char_start=np.asarray([word["char_start"] for word in words], dtype=np.int32),
                        word_char_end=np.asarray([word["char_end"] for word in words], dtype=np.int32),
                        text_word_feat=info["features"].astype(np.float32),
                        text_word_valid=info["valid"].astype(np.uint8),
                        roberta_token_ids=info["token_ids"].astype(np.int32),
                        roberta_token_offsets=info["token_offsets"].astype(np.int32),
                        roberta_token_word_index=info["token_word_index"].astype(np.int32),
                        roberta_word_token_indptr=info["word_token_indptr"].astype(np.int32),
                        roberta_word_token_indices=info["word_token_indices"].astype(np.int32),
                    )
                temp.replace(temp_path)
                record.update({"status": "PASS", "word_count": len(words), "token_count": info["input_length"], "text_feature_shape": list(info["features"].shape), "text_cache_sha256": sha256(temp_path), "roberta_model_sha256": model_sha, "roberta_asset_identity_sha256": model_identity_sha, "config_sha256": cfg_sha})
        except core.HardStop as exc:
            record.update({"status": "HARD_STOP", "error": f"{type(exc).__name__}: {exc}"})
            hard_stop = True
        except Exception as exc:
            record.update({"status": "SAMPLE_FAIL", "error": f"{type(exc).__name__}: {exc}"})
        record["runtime_sec"] = time.perf_counter() - started
        write_json(log_path, record)
        results.append(record)
        print(f"TEXT {index}/{len(rows)} {key} {record['status']}", flush=True)
        if hard_stop:
            break
    write_json(output_root / "reports" / "text_stage_summary.json", {"sample_count": len(rows), "pass_count": sum(x["status"] in {"PASS", "RESUMED_TEXT_CACHE"} for x in results), "failures": [x for x in results if x["status"] == "SAMPLE_FAIL"], "results": results})
    if hard_stop:
        return 2
    return 0 if all(x["status"] in {"PASS", "RESUMED_TEXT_CACHE"} for x in results) else 1


def alignment_route(row: dict[str, Any], core, audio_end: float, video_start: float, video_end: float):
    audit = row["audit_row"]
    speech_state = str(audit.get("audio_speech_valid", "unknown"))
    correspondence = str(audit.get("text_audio_correspondence", "unknown"))
    status = str(audit.get("text_av_time_mapping_status", "unverified"))
    source_words = core.source_words(row["official_text"])
    n = len(source_words)
    t0 = row["shared_t0_sec"]
    word_start = np.full(n, np.nan, dtype=np.float64)
    word_end = np.full(n, np.nan, dtype=np.float64)
    word_time_valid = np.zeros(n, dtype=np.uint8)
    trace_words = row["alignment_trace"].get("mapped_words", [])
    use_word_times = correspondence == "confirmed_match" and status in {"word_valid", "word_partial"}
    if use_word_times and len(trace_words) == n:
        previous_start = -np.inf
        previous_end = -np.inf
        for wi, mapped in enumerate(trace_words):
            if int(mapped.get("word_index", -1)) != wi or str(mapped.get("word_text", "")) != source_words[wi]["text"]:
                raise PipelineStop(f"word alignment identity mismatch for {row['sample_key']} at word {wi}")
            start, end = mapped.get("start_sec"), mapped.get("end_sec")
            legal = bool(mapped.get("alignment_valid", 0)) and start is not None and end is not None
            if legal:
                start, end = float(start), float(end)
                legal = np.isfinite(start) and np.isfinite(end) and end > start
                legal = legal and start >= max(t0, video_start) - 1e-6 and end <= min(audio_end, video_end) + 1e-6
                legal = legal and start >= previous_start - 1e-6 and end >= previous_end - 1e-6
            if legal:
                word_start[wi], word_end[wi], word_time_valid[wi] = start, end, 1
                previous_start, previous_end = start, end
    elif use_word_times and len(trace_words) != n:
        use_word_times = False
    has_word_mapping = bool(word_time_valid.any())
    if correspondence == "confirmed_match" and has_word_mapping:
        mode = "TRI_MODAL_WORD_VALID"
        granularity = "word"
        time_status = "word_valid" if bool(word_time_valid.all()) else "word_partial"
        evidence = "existing_audit_human_confirmed_match_plus_hash_verified_stable_ts_word_trace"
        speech_valid = 1
    elif speech_state == "0" or correspondence == "no_speech":
        mode = "TRI_MODAL_WITH_AUDIO_CONTENT_INVALID"
        granularity = "clip"
        time_status = "clip_only"
        evidence = "full100_audit_confirmed_digital_silence_or_no_speech"
        speech_valid = 0
    elif correspondence == "confirmed_mismatch":
        mode = "AV_VALID_TEXT_UNALIGNED"
        granularity = "clip"
        time_status = "unavailable"
        evidence = "existing_audit_human_confirmed_transcript_audio_mismatch"
        speech_valid = 1
    elif correspondence == "confirmed_match":
        mode = "UNCERTAIN_REVIEW"
        granularity = "clip"
        time_status = "clip_only"
        evidence = "content_match_evidence_exists_but_no_usable_word_trace; conservative_clip_only_output"
        speech_valid = 1
    elif not has_word_mapping:
        mode = "UNCERTAIN_REVIEW"
        granularity = "clip"
        time_status = "clip_only"
        evidence = "machine_screen_only_no_content_assertion; conservative_clip_only_output"
        speech_valid = -1
    else:
        raise PipelineStop(f"unsupported route for {row['sample_key']}")
    if not has_word_mapping:
        word_start[:] = np.nan
        word_end[:] = np.nan
        word_time_valid[:] = 0
    return {
        "words": source_words,
        "word_start": word_start,
        "word_end": word_end,
        "word_time_valid": word_time_valid,
        "alignment_mode": mode,
        "alignment_granularity": granularity,
        "text_audio_correspondence": correspondence if correspondence in {"confirmed_match", "confirmed_mismatch", "no_speech"} else "not_asserted",
        "text_av_time_mapping_status": time_status,
        "audio_speech_valid": speech_valid,
        "audio_speech_evidence": evidence,
    }


def media_stage(args: argparse.Namespace) -> int:
    import gc
    import psutil
    import mediapipe as mp
    core = import_core()
    output_root = Path(args.output_root).resolve()
    rows = selected_rows(output_root, args.sample_key)
    verify_live_label(Path(args.input_root), rows)
    asset_root = Path(args.asset_root)
    asset_info = check_assets(asset_root)
    saved_assets = json.loads((output_root / "metadata" / "model_assets.json").read_text(encoding="utf-8"))
    if asset_info != saved_assets:
        raise PipelineStop("model assets changed after prepare")
    face_path = asset_root / "mediapipe" / "face_landmarker.task"
    options = mp.tasks.vision.FaceLandmarkerOptions(
        base_options=mp.tasks.BaseOptions(model_asset_path=str(face_path)),
        running_mode=mp.tasks.vision.RunningMode.VIDEO,
        num_faces=1,
        output_face_blendshapes=True,
    )
    manifest_rows: list[dict[str, Any]] = []
    hard_stop = None
    for index, row in enumerate(rows, 1):
        key, stub = row["sample_key"], safe_stem(row["sample_key"])
        text_cache = output_root / ".cache" / "text" / f"{stub}.npz"
        output_file = output_root / "features" / f"{stub}.npz"
        log_file = output_root / "logs" / f"{stub}.media.json"
        started = time.perf_counter()
        record: dict[str, Any] = {"sample_key": key, "stage": "media", "status": "RUNNING", "source_mp4_sha256": row["source_mp4_sha256"]}
        landmarker = None
        try:
            if not text_cache.is_file():
                raise RuntimeError(f"text feature cache missing for {key}; text stage did not complete")
            text_log_path = output_root / "logs" / f"{stub}.text.json"
            if not text_log_path.is_file():
                raise RuntimeError(f"text cache sidecar is missing for {key}")
            text_log = json.loads(text_log_path.read_text(encoding="utf-8"))
            cache_sha = sha256(text_cache)
            if text_log.get("text_cache_sha256") != cache_sha:
                raise RuntimeError(f"text cache hash differs from its sidecar for {key}")
            model_sha = sha256(asset_root / "roberta-base" / ROBERTA_REVISION / "model.safetensors")
            text_arrays = validate_text_cache(text_cache, row, core, model_sha, asset_info["roberta_asset_identity_sha256"], config_digest(output_root))
            words = [{"text": str(text), "char_start": int(start), "char_end": int(end)} for text, start, end in zip(text_arrays["words"], text_arrays["word_char_start"], text_arrays["word_char_end"])]
            if len(words) != len(text_arrays["text_word_feat"]) or len(words) == 0:
                raise PipelineStop(f"text cache shape is inconsistent for {key}")
            video_path = Path(args.input_root) / row["source_relpath"]
            if not video_path.is_file() or sha256(video_path) != row["source_mp4_sha256"]:
                raise PipelineStop(f"source MP4 changed after preflight for {key}")
            shared_t0 = float(row["shared_t0_sec"])
            audio16, audio_info = core.decode_audio_16k(video_path, shared_t0)
            if not np.isfinite(audio16).all() or audio16.size == 0:
                raise PipelineStop(f"decoded source audio is empty/nonfinite for {key}")
            lld = core.opensmile_features(audio16)
            if lld["config_sha256"] != asset_info["opensmile_config_sha256"] or lld["feature_names"] != asset_info["opensmile_feature_names"]:
                raise PipelineStop(f"openSMILE config or feature order changed for {key}")
            if not np.isfinite(lld["values"]).all():
                raise RuntimeError(f"openSMILE returned nonfinite LLD values for {key}")
            if lld["values"].shape[1] != 25:
                raise PipelineStop(f"openSMILE LLD dimension changed for {key}")

            landmarker = mp.tasks.vision.FaceLandmarker.create_from_options(options)
            video_info, video_pts, video_support_start, video_support_end, blend, face_valid = core.decode_video_blendshapes(video_path, shared_t0, landmarker)
            landmarker.close()
            landmarker = None
            common_start = max(float(audio_info["resampled_start_sec"]), float(video_info["actual_video_coverage_start_sec"]))
            common_end = min(float(audio_info["resampled_end_sec"]), float(video_info["actual_video_coverage_end_sec"]))
            if common_end <= common_start:
                raise PipelineStop(f"decoded audio/video have no positive shared presentation coverage for {key}")
            route = alignment_route(row, core, float(audio_info["resampled_end_sec"]), common_start, float(video_info["actual_video_coverage_end_sec"]))
            rel_word_start = route["word_start"] - float(audio_info["resampled_start_sec"])
            rel_word_end = route["word_end"] - float(audio_info["resampled_start_sec"])
            word_audio, audio_word_valid, audio_indptr, audio_indices, audio_issues = core.aggregate_by_center(
                route["words"], rel_word_start, rel_word_end, lld["centers"], lld["values"], dim=25
            )
            audio_word_valid &= route["word_time_valid"]
            word_audio[audio_word_valid == 0] = 0.0
            word_vision, vision_word_valid, vision_indptr, vision_indices, vision_issues = core.aggregate_by_center(
                route["words"], route["word_start"] - shared_t0, route["word_end"] - shared_t0,
                video_pts - shared_t0, blend, valid_frames=face_valid, dim=52,
            )
            vision_word_valid &= route["word_time_valid"]
            word_vision[vision_word_valid == 0] = 0.0
            text_valid = text_arrays["text_word_valid"].astype(np.uint8)
            text_feat = text_arrays["text_word_feat"].astype(np.float32)
            text_feat[text_valid == 0] = 0.0
            words_text = [word["text"] for word in route["words"]]
            npz_arrays: dict[str, Any] = {
                "schema_version": np.asarray(SCHEMA_VERSION),
                "config_sha256": np.asarray(config_digest(output_root)),
                "sample_key": np.asarray(key),
                "video_id": np.asarray(row["video_id"]),
                "clip_id": np.asarray(row["clip_id"]),
                "official_text": np.asarray(row["official_text"]),
                "source_video_sha256": np.asarray(row["source_mp4_sha256"]),
                "source_label_xlsx_sha256": np.asarray(row["source_label_xlsx_sha256"]),
                "roberta_model_sha256": np.asarray(asset_info["roberta_model_sha256"]),
                "roberta_asset_identity_sha256": np.asarray(asset_info["roberta_asset_identity_sha256"]),
                "mediapipe_model_sha256": np.asarray(asset_info["mediapipe_model_sha256"]),
                "alignment_trace_sha256": np.asarray(row["alignment_trace_sha256"]),
                "alignment_mode": np.asarray(route["alignment_mode"]),
                "alignment_granularity": np.asarray(route["alignment_granularity"]),
                "text_audio_correspondence": np.asarray(route["text_audio_correspondence"]),
                "text_av_time_mapping_status": np.asarray(route["text_av_time_mapping_status"]),
                "audio_speech_evidence": np.asarray(route["audio_speech_evidence"]),
                "text_present": np.asarray(1, dtype=np.uint8),
                "text_content_valid": np.asarray(1, dtype=np.uint8),
                "audio_present": np.asarray(1, dtype=np.uint8),
                "audio_observation_valid": np.asarray(1, dtype=np.uint8),
                "audio_speech_valid": np.asarray(route["audio_speech_valid"], dtype=np.int8),
                "video_present": np.asarray(1, dtype=np.uint8),
                "audio_visual_time_valid": np.asarray(1, dtype=np.uint8),
                "vision_feature_valid": np.asarray(int(face_valid.any()), dtype=np.uint8),
                "face_feature_valid": np.asarray(int(face_valid.any()), dtype=np.uint8),
                "words": np.asarray(words_text, dtype=np.str_),
                "word_char_start": np.asarray([word["char_start"] for word in route["words"]], dtype=np.int32),
                "word_char_end": np.asarray([word["char_end"] for word in route["words"]], dtype=np.int32),
                "word_start_sec": route["word_start"].astype(np.float64),
                "word_end_sec": route["word_end"].astype(np.float64),
                "word_time_valid": route["word_time_valid"].astype(np.uint8),
                "text_word_feat": text_feat,
                "text_word_valid": text_valid,
                "word_audio_feat": word_audio.astype(np.float32),
                "word_audio_valid": audio_word_valid.astype(np.uint8),
                "word_vision_feat": word_vision.astype(np.float32),
                "word_vision_valid": vision_word_valid.astype(np.uint8),
                "raw_audio_lld_values": lld["values"].astype(np.float32),
                "raw_audio_lld_start_sec": lld["starts"].astype(np.float64) + float(audio_info["resampled_start_sec"]),
                "raw_audio_lld_end_sec": lld["ends"].astype(np.float64) + float(audio_info["resampled_start_sec"]),
                "raw_audio_lld_center_sec": lld["centers"].astype(np.float64) + float(audio_info["resampled_start_sec"]),
                "raw_audio_lld_valid": np.ones(len(lld["values"]), dtype=np.uint8),
                "raw_audio_lld_feature_names": np.asarray(lld["feature_names"], dtype=np.str_),
                "opensmile_config_sha256": np.asarray(lld["config_sha256"]),
                "raw_video_blendshape_values": blend.astype(np.float32),
                "raw_video_pts_sec": video_pts.astype(np.float64),
                "raw_video_support_start_sec": video_support_start.astype(np.float64),
                "raw_video_support_end_sec": video_support_end.astype(np.float64),
                "raw_video_face_feature_valid": face_valid.astype(np.uint8),
                "raw_video_blendshape_names": np.asarray(core.BLENDSHAPE_NAMES, dtype=np.str_),
                "audio_word_feature_indptr": audio_indptr.astype(np.int32),
                "audio_word_feature_indices": audio_indices.astype(np.int32),
                "vision_word_feature_indptr": vision_indptr.astype(np.int32),
                "vision_word_feature_indices": vision_indices.astype(np.int32),
                "roberta_token_ids": text_arrays["roberta_token_ids"],
                "roberta_token_offsets": text_arrays["roberta_token_offsets"],
                "roberta_token_word_index": text_arrays["roberta_token_word_index"],
                "roberta_word_token_indptr": text_arrays["roberta_word_token_indptr"],
                "roberta_word_token_indices": text_arrays["roberta_word_token_indices"],
                "audio_presentation_coverage_sec": np.asarray([audio_info["resampled_start_sec"], audio_info["resampled_end_sec"]], dtype=np.float64),
                "video_presentation_coverage_sec": np.asarray([video_info["actual_video_coverage_start_sec"], video_info["actual_video_coverage_end_sec"]], dtype=np.float64),
                "av_common_coverage_sec": np.asarray([common_start, common_end], dtype=np.float64),
                "shared_t0_sec": np.asarray(shared_t0, dtype=np.float64),
                "audio_observation_length": np.asarray(len(lld["values"]), dtype=np.int32),
                "video_observation_length": np.asarray(len(video_pts), dtype=np.int32),
                "text_sequence_length": np.asarray(len(words_text), dtype=np.int32),
                "original_effective_duration_sec": np.asarray(max(audio_info["resampled_end_sec"], video_info["actual_video_coverage_end_sec"]) - shared_t0, dtype=np.float64),
                "alignment_zero_duration_word_count": np.asarray(int(row["audit_row"].get("zero_duration_word_count", 0)), dtype=np.int32),
                "alignment_mapping_fail_count": np.asarray(int(row["audit_row"].get("mapping_fail_count", 0)), dtype=np.int32),
                "diagnostic_alignment_probability_summary_json": np.asarray(row["audit_row"].get("alignment_probability_summary", "{}")),
            }
            temp = output_file.with_suffix(".npz.tmp")
            with temp.open("wb") as stream:
                np.savez_compressed(stream, **npz_arrays)
            temp.replace(output_file)
            record.update({
                "status": "PASS",
                "alignment_mode": route["alignment_mode"],
                "alignment_granularity": route["alignment_granularity"],
                "text_av_time_mapping_status": route["text_av_time_mapping_status"],
                "text_audio_correspondence": route["text_audio_correspondence"],
                "audio_speech_valid": int(route["audio_speech_valid"]),
                "audio_speech_evidence": route["audio_speech_evidence"],
                "official_word_count": len(words_text),
                "valid_word_time_count": int(route["word_time_valid"].sum()),
                "text_word_valid_count": int(text_valid.sum()),
                "audio_word_valid_count": int(audio_word_valid.sum()),
                "vision_word_valid_count": int(vision_word_valid.sum()),
                "audio_lld_frame_count": len(lld["values"]),
                "video_frame_count": len(video_pts),
                "face_detection_ratio": float(face_valid.mean()) if len(face_valid) else 0.0,
                "audio_coverage_sec": [float(audio_info["resampled_start_sec"]), float(audio_info["resampled_end_sec"])],
                "video_coverage_sec": [float(video_info["actual_video_coverage_start_sec"]), float(video_info["actual_video_coverage_end_sec"])],
                "audio_word_issues": audio_issues,
                "vision_word_issues": vision_issues,
                "output_file": str(output_file.relative_to(output_root)),
                "output_bytes": output_file.stat().st_size,
                "output_sha256": sha256(output_file),
                "process_rss_bytes": int(psutil.Process().memory_info().rss),
                "config_sha256": config_digest(output_root),
                "opensmile_config_sha256": lld["config_sha256"],
                "opensmile_feature_names": lld["feature_names"],
                "roberta_asset_identity_sha256": asset_info["roberta_asset_identity_sha256"],
                "mediapipe_model_sha256": asset_info["mediapipe_model_sha256"],
            })
        except core.HardStop as exc:
            record.update({
                "status": "HARD_STOP",
                "anomaly_type": "EXTRACTION_OR_TIMELINE_ANOMALY",
                "error": f"{type(exc).__name__}: {exc}",
            })
            hard_stop = str(exc)
        except PipelineStop as exc:
            message = str(exc)
            lowered = message.casefold()
            anomaly_type = "EXTRACTION_OR_TIMELINE_ANOMALY" if any(token in lowered for token in ("source mp4", "timeline", "presentation", "pts", "decode", "audio/video")) else "PIPELINE_CONTRACT_HARD_STOP"
            record.update({"status": "HARD_STOP", "anomaly_type": anomaly_type, "error": message})
            hard_stop = str(exc)
        except Exception as exc:
            record.update({"status": "SAMPLE_FAIL", "error": f"{type(exc).__name__}: {exc}"})
        finally:
            if landmarker is not None:
                try:
                    landmarker.close()
                except Exception:
                    pass
            record["runtime_sec"] = time.perf_counter() - started
            write_json(log_file, record)
            if record.get("status") == "PASS":
                manifest_rows.append({
                    "sample_key": key,
                    "source_mp4_sha256": row["source_mp4_sha256"],
                    "official_text_sha256": row["official_text_sha256"],
                    "output_file": record["output_file"],
                    "output_sha256": record["output_sha256"],
                    "output_bytes": record["output_bytes"],
                    "alignment_mode": record["alignment_mode"],
                    "alignment_granularity": record["alignment_granularity"],
                    "text_audio_correspondence": record.get("text_audio_correspondence", "not_asserted"),
                    "text_av_time_mapping_status": record["text_av_time_mapping_status"],
                    "audio_present": 1,
                    "audio_observation_valid": 1,
                    "audio_speech_valid": record.get("audio_speech_valid", -1),
                    "video_present": 1,
                    "face_feature_valid": int(record.get("face_detection_ratio", 0) > 0),
                    "audio_visual_time_valid": 1,
                    "official_word_count": record.get("official_word_count"),
                    "valid_word_time_count": record.get("valid_word_time_count"),
                    "text_word_valid_count": record.get("text_word_valid_count"),
                    "audio_word_valid_count": record.get("audio_word_valid_count"),
                    "vision_word_valid_count": record.get("vision_word_valid_count"),
                    "audio_lld_frame_count": record.get("audio_lld_frame_count"),
                    "video_frame_count": record.get("video_frame_count"),
                    "face_detection_ratio": record.get("face_detection_ratio"),
                    "runtime_sec": record.get("runtime_sec"),
                    "status": record.get("status"),
                })
            print(f"MEDIA {index}/{len(rows)} {key} {record['status']}", flush=True)
        if hard_stop:
            break
        gc.collect()
    if manifest_rows:
        fields = list(manifest_rows[0])
        write_csv(output_root / "manifest.csv", manifest_rows, fields)
    write_json(output_root / "reports" / "media_stage_summary.json", {
        "requested_count": len(rows),
        "written_count": len(manifest_rows),
        "pass_count": sum(row["status"] == "PASS" for row in manifest_rows),
        "hard_stop": hard_stop,
        "sample_failures": [json.loads(path.read_text(encoding="utf-8")) for path in sorted((output_root / "logs").glob("*.media.json")) if json.loads(path.read_text(encoding="utf-8")).get("status") == "SAMPLE_FAIL"],
    })
    return 2 if hard_stop else (0 if len(manifest_rows) == len(rows) else 1)


def run_child(args: argparse.Namespace, stage: str, *, sample_key: str | None = None, timeout_sec: int = 3600, max_attempts: int = 2, retry_resume: bool = False) -> int:
    output_root = Path(args.output_root).resolve()
    log_root = output_root / "logs"
    attempts_root = log_root / "attempts" / (safe_stem(sample_key) if sample_key else stage)
    attempts_root.mkdir(parents=True, exist_ok=True)
    final_rc = 1
    for attempt in range(1, max_attempts + 1):
        cmd = [sys.executable, str(Path(__file__).resolve()), "--stage", stage,
               "--input-root", str(args.input_root), "--audit-root", str(args.audit_root),
               "--asset-root", str(args.asset_root), "--output-root", str(args.output_root)]
        if sample_key:
            cmd.append(f"--sample-key={sample_key}")
        elif args.sample_key:
            for key in args.sample_key:
                cmd.append(f"--sample-key={key}")
        if args.resume or (retry_resume and attempt > 1):
            cmd.append("--resume")
        try:
            child = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout_sec)
            stdout, stderr, final_rc = child.stdout, child.stderr, child.returncode
            state = "EXITED"
        except subprocess.TimeoutExpired as exc:
            stdout = exc.stdout.decode("utf-8", "replace") if isinstance(exc.stdout, bytes) else (exc.stdout or "")
            stderr = exc.stderr.decode("utf-8", "replace") if isinstance(exc.stderr, bytes) else (exc.stderr or "")
            stderr += f"\nTIMEOUT after {timeout_sec} seconds\n"
            final_rc, state = 124, "TIMEOUT"
        prefix = attempts_root / f"{stage}_attempt_{attempt}"
        prefix.with_suffix(".stdout.log").write_text(stdout, encoding="utf-8")
        prefix.with_suffix(".stderr.log").write_text(stderr, encoding="utf-8")
        prefix.with_suffix(".exit_code.txt").write_text(f"{final_rc}\n", encoding="ascii")
        write_json(prefix.with_suffix(".json"), {
            "stage": stage, "sample_key": sample_key, "attempt": attempt, "state": state,
            "exit_code": final_rc, "timeout_sec": timeout_sec, "same_parameters_retry": attempt > 1,
        })
        print(stdout[-3000:], end="", flush=True)
        if final_rc == 0:
            return 0
        if final_rc == 2:
            if sample_key:
                event_path = log_root / f"{safe_stem(sample_key)}.{stage}.json"
                prior = {}
                if event_path.is_file():
                    try:
                        prior = json.loads(event_path.read_text(encoding="utf-8"))
                    except Exception:
                        prior = {}
                prior.update({
                    "sample_key": sample_key, "stage": stage, "status": "HARD_STOP",
                    "process_exit_code": final_rc,
                    "error": prior.get("error") or stderr[-3000:] or "worker returned HARD_STOP exit code 2",
                })
                write_json(event_path, prior)
            print(stderr[-3000:], file=sys.stderr, flush=True)
            return 2
        if attempt < max_attempts:
            print(f"RETRY {stage} {sample_key or ''} after exit={final_rc}; same parameters", flush=True)
    if sample_key:
        log_path = log_root / f"{safe_stem(sample_key)}.{stage}.json"
        previous = {}
        if log_path.is_file():
            try:
                previous = json.loads(log_path.read_text(encoding="utf-8"))
            except Exception:
                previous = {}
        previous.update({
            "sample_key": sample_key, "stage": stage,
            "status": "HARD_STOP" if final_rc == 2 else "SAMPLE_FAIL",
            "process_exit_code": final_rc,
            "error": previous.get("error") or ("worker exceeded fixed timeout" if final_rc == 124 else f"worker exited {final_rc} after {max_attempts} attempts"),
        })
        write_json(log_path, previous)
    print(stderr[-3000:], file=sys.stderr, flush=True)
    return final_rc


def media_resume_valid(output_root: Path, row: dict[str, Any]) -> bool:
    stub = safe_stem(row["sample_key"])
    log_path = output_root / "logs" / f"{stub}.media.json"
    feature_path = output_root / "features" / f"{stub}.npz"
    if not log_path.is_file() or not feature_path.is_file():
        return False
    try:
        log = json.loads(log_path.read_text(encoding="utf-8"))
        if log.get("status") != "PASS" or log.get("output_sha256") != sha256(feature_path):
            return False
        if log.get("source_mp4_sha256") != row["source_mp4_sha256"] or log.get("config_sha256") != config_digest(output_root):
            return False
        with np.load(feature_path, allow_pickle=False) as z:
            return (str(z["sample_key"].item()) == row["sample_key"]
                    and str(z["source_video_sha256"].item()) == row["source_mp4_sha256"]
                    and str(z["config_sha256"].item()) == config_digest(output_root))
    except Exception:
        return False


def all_stages(args: argparse.Namespace) -> int:
    selected = prepare(args)
    output_root = Path(args.output_root).resolve()
    import psutil
    memory_floor = 1024**3
    text_hard_stop = False
    text_paused = False
    for index, row in enumerate(selected, 1):
        key = row["sample_key"]
        text_log = output_root / "logs" / f"{safe_stem(key)}.text.json"
        text_cache = output_root / ".cache" / "text" / f"{safe_stem(key)}.npz"
        if args.resume and text_log.is_file() and text_cache.is_file():
            try:
                item = json.loads(text_log.read_text(encoding="utf-8"))
                if item.get("status") in {"PASS", "RESUMED_TEXT_CACHE"} and item.get("text_cache_sha256") == sha256(text_cache):
                    print(f"TEXT {index}/{len(selected)} {key} RESUMED_VALID_CACHE", flush=True)
                    continue
            except Exception:
                pass
        available = int(psutil.virtual_memory().available)
        if available < memory_floor:
            write_json(text_log, {
                "sample_key": key, "stage": "text", "status": "PAUSED_RESOURCE",
                "available_ram_bytes": available, "memory_floor_bytes": memory_floor,
                "error": "available memory below 1 GiB before sample boundary",
            })
            text_paused = True
            print(f"PAUSED_RESOURCE before text sample {key}: available_ram_bytes={available}", flush=True)
            continue
        print(f"TEXT_SAMPLE_START {index}/{len(selected)} {key}; available_ram_bytes={available}", flush=True)
        rc = run_child(args, "text", sample_key=key, timeout_sec=300, max_attempts=2, retry_resume=True)
        if rc == 2:
            text_hard_stop = True
            break

    text_logs = []
    for row in selected:
        path = output_root / "logs" / f"{safe_stem(row['sample_key'])}.text.json"
        if path.is_file():
            text_logs.append(json.loads(path.read_text(encoding="utf-8")))
    write_json(output_root / "reports" / "text_stage_summary.json", {
        "requested_count": len(selected),
        "pass_count": sum(item.get("status") in {"PASS", "RESUMED_TEXT_CACHE"} for item in text_logs),
        "failure_count": sum(item.get("status") not in {"PASS", "RESUMED_TEXT_CACHE"} for item in text_logs),
        "paused_for_resource": text_paused, "hard_stop": text_hard_stop,
        "sample_results": text_logs,
    })
    if text_hard_stop:
        for row in selected:
            log_path = output_root / "logs" / f"{safe_stem(row['sample_key'])}.media.json"
            if not args.resume or not media_resume_valid(output_root, row):
                write_json(log_path, {"sample_key": row["sample_key"], "stage": "media", "status": "NOT_RUN_AFTER_HARD_STOP", "source_mp4_sha256": row["source_mp4_sha256"]})
    elif text_paused:
        for row in selected:
            log_path = output_root / "logs" / f"{safe_stem(row['sample_key'])}.media.json"
            if not args.resume or not media_resume_valid(output_root, row):
                write_json(log_path, {"sample_key": row["sample_key"], "stage": "media", "status": "PAUSED_RESOURCE", "source_mp4_sha256": row["source_mp4_sha256"], "error": "text stage paused at resource boundary"})
    else:
        media_hard_stop = False
        media_resource_pause = False
        for index, row in enumerate(selected, 1):
            key = row["sample_key"]
            if args.resume and media_resume_valid(output_root, row):
                print(f"MEDIA {index}/{len(selected)} {key} RESUMED_VALID_OUTPUT", flush=True)
                continue
            available = int(psutil.virtual_memory().available)
            if available < memory_floor:
                record = {
                    "sample_key": key, "stage": "media", "status": "PAUSED_RESOURCE",
                    "source_mp4_sha256": row["source_mp4_sha256"], "available_ram_bytes": available,
                    "memory_floor_bytes": memory_floor, "error": "available memory below 1 GiB before sample boundary",
                }
                write_json(output_root / "logs" / f"{safe_stem(key)}.media.json", record)
                media_resource_pause = True
                for pending in selected[index:]:
                    pending_path = output_root / "logs" / f"{safe_stem(pending['sample_key'])}.media.json"
                    if not args.resume or not media_resume_valid(output_root, pending):
                        write_json(pending_path, {
                            "sample_key": pending["sample_key"], "stage": "media", "status": "PAUSED_RESOURCE",
                            "source_mp4_sha256": pending["source_mp4_sha256"], "memory_floor_bytes": memory_floor,
                            "error": "not started after previous sample-boundary resource pause",
                        })
                print(f"PAUSED_RESOURCE before media sample {key}: available_ram_bytes={available}", flush=True)
                break
            print(f"MEDIA_SAMPLE_START {index}/{len(selected)} {key}; available_ram_bytes={available}", flush=True)
            rc = run_child(args, "media", sample_key=key, timeout_sec=900, max_attempts=2)
            if rc == 2:
                media_hard_stop = True
                for pending in selected[index:]:
                    pending_path = output_root / "logs" / f"{safe_stem(pending['sample_key'])}.media.json"
                    if not args.resume or not media_resume_valid(output_root, pending):
                        write_json(pending_path, {
                            "sample_key": pending["sample_key"], "stage": "media", "status": "NOT_RUN_AFTER_HARD_STOP",
                            "source_mp4_sha256": pending["source_mp4_sha256"], "error": f"stopped after hard stop at {key}",
                        })
                break

    manifest_rows = rebuild_output_tables(output_root, selected)
    write_json(output_root / "reports" / "media_stage_summary.json", {
        "requested_count": len(selected),
        "pass_count": sum(row["status"] == "PASS" for row in manifest_rows),
        "failure_count": sum(row["status"] != "PASS" for row in manifest_rows),
        "status_counts": {status: sum(row["status"] == status for row in manifest_rows) for status in sorted({row["status"] for row in manifest_rows})},
        "sample_failures": [{"sample_key": row["sample_key"], "status": row["status"], "error": row["error"]} for row in manifest_rows if row["status"] != "PASS"],
    })
    validator = Path(__file__).resolve().with_name("validate_q1_v1.py")
    cmd = [sys.executable, str(validator), "--output-root", str(output_root), "--input-root", str(args.input_root),
           "--audit-root", str(args.audit_root), "--asset-root", str(args.asset_root), "--expected-count", str(len(selected))]
    try:
        check = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=1800)
        check_rc, check_stdout, check_stderr = check.returncode, check.stdout, check.stderr
    except subprocess.TimeoutExpired as exc:
        check_rc = 124
        check_stdout = exc.stdout.decode("utf-8", "replace") if isinstance(exc.stdout, bytes) else (exc.stdout or "")
        check_stderr = (exc.stderr.decode("utf-8", "replace") if isinstance(exc.stderr, bytes) else (exc.stderr or "")) + "\nVALIDATOR_TIMEOUT after 1800 seconds\n"
    log_root = output_root / "logs"
    (log_root / "validator.stdout.log").write_text(check_stdout, encoding="utf-8")
    (log_root / "validator.stderr.log").write_text(check_stderr, encoding="utf-8")
    (log_root / "validator.exit_code.txt").write_text(f"{check_rc}\n", encoding="ascii")
    print(check_stdout, end="", flush=True)
    if check_rc:
        print(check_stderr, file=sys.stderr, flush=True)
    if text_hard_stop or ("media_hard_stop" in locals() and media_hard_stop):
        return 2
    if text_paused or ("media_resource_pause" in locals() and media_resource_pause):
        return 1
    if check_rc:
        return check_rc
    if any(row["status"] != "PASS" for row in manifest_rows):
        return 1
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Q1 full100 multimodal feature production, schema v1.0")
    parser.add_argument("--stage", choices=("prepare", "text", "media", "all"), default="all")
    parser.add_argument("--input-root", type=Path, default=DEFAULT_INPUT_ROOT)
    parser.add_argument("--audit-root", type=Path, default=DEFAULT_AUDIT_ROOT)
    parser.add_argument("--asset-root", type=Path, default=DEFAULT_ASSET_ROOT)
    parser.add_argument("--output-root", type=Path, default=DELIVERY_ROOT)
    parser.add_argument("--sample-key", action="append", default=[])
    parser.add_argument("--resume", action="store_true", help="resume only with matching frozen config, source/model hashes, and validated cache/output SHA sidecars")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if args.stage == "prepare":
            prepare(args)
            print("PREPARE PASS", flush=True)
            return 0
        if args.stage == "text":
            return text_stage(args)
        if args.stage == "media":
            return media_stage(args)
        return all_stages(args)
    except PipelineStop as exc:
        print(f"HARD_STOP: {exc}", file=sys.stderr, flush=True)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
